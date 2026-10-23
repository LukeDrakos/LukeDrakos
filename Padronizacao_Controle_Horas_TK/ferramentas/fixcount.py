import openpyxl,sys
wb=openpyxl.load_workbook(sys.argv[1])
lh=wb['Lanc_HORAS']; n=lh.max_row
print('LH',lh['AS2'].value,lh['AS3'].value,lh['AU3'].value, n)
for r in range(2,n+1):
    lh[f'AS{r}']=f'=IF($AR{r}=1,COUNTIF($AR$2:$AR{r},1),0)'
    lh[f'AU{r}']=f'=IF($AT{r}=1,COUNTIF($AT$2:$AT{r},1),0)'
he=wb['Lanc_HE']; print('HE',he['AG2'].value,he['AG3'].value)
for r in range(2,he.max_row+1):
    if he[f'AF{r}'].value: he[f'AG{r}']=f'=IF($AF{r}=1,COUNTIF($AF$2:$AF{r},1),0)'
cp=wb['Compensacoes']; print('CP',cp['O4'].value,cp['O5'].value)
for r in range(4,cp.max_row+1):
    if cp[f'N{r}'].value: cp[f'O{r}']=f'=IF($N{r}=1,COUNTIF($N$4:$N{r},1),0)'
wb.save(sys.argv[2]); print('ok')
