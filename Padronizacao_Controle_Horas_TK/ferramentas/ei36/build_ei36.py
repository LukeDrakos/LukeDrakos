"""EI36 v3 — lógica da EI32 (v55) aplicada na EI36 v2 do Felipe, mantendo o layout dele.

Uso: python build_ei36.py <EI36 v2.xlsx> <EI32 v55.xlsx> <Pontos 2026.xlsx> <saida.xlsx>

O que faz (tudo valendo desde 01/01/2026, decisão do Lucas em 23/10):
  - Lancamentos: cascata do ferista (férias primeiro, doação mínima 1h), compensação em bloco no mês,
    abatimento de Valor/Função, Turno 2 automático, colunas da macro (divergência / conferido), relatório.
  - Tabelas: 6 status novos, parâmetros novos, tratamentos de HE, calendário PCAM, nomes da macro.
  - Compensacoes: abate em bloco (só "Aprovado? = Sim"); falta vinculada vira opcional.
  - Resumo_Mensal / Painel: compensação em bloco + Valor/Função; total e "confere" atualizados.
  - Abas novas: Relatorio (PDF) e Conferencia_Ponto (macro).
  - Cadastro: RG provisório 9000xx trocado pela matrícula do relógio (mesmo nome na base de pontos).
  - Visual: paleta da EI32 (azul TK / marinho / cinza), sem mudar posição de nada.
"""
import re, sys, unicodedata, datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

SRC, E32, PONTOS, OUT = sys.argv[1:5]
wb = openpyxl.load_workbook(SRC)
LAST = 3489                      # última linha de Lancamentos (faixas dos nomes)
HLAST = 1000                     # última linha de Lanc_HE
CLAST = 999                      # última linha de Compensacoes

TK, NAVY, TXT, CINZA, BORDA = 'FF0095D8', 'FF1F3864', 'FF4D4D4D', 'FF7F7F7F', 'FFD9D9D9'
VERDE, VERM, F2, D9 = 'FFD5F5E3', 'FFFADBD8', 'FFF2F2F2', 'FFD9D9D9'
fill = lambda c: PatternFill('solid', fgColor=c)
fino = Side(style='thin', color=BORDA)
BOX = Border(left=fino, right=fino, top=fino, bottom=fino)
branco = Side(style='thin', color='FFFFFFFF')
BOXH = Border(left=branco, right=branco, top=branco, bottom=branco)
HDR_FONT = Font(name='Arial', size=9, bold=True, color='FFFFFFFF')
HDR_AL = Alignment(horizontal='center', vertical='center', wrap_text=True)


def nome(n, ref):
    if n in wb.defined_names:
        del wb.defined_names[n]
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def tpl(formula, row=4):
    """fórmula da linha `row` -> modelo com {r} (só referências relativas de linha)."""
    f = formula.replace('{', '{{').replace('}', '}}')
    return re.sub(r'(?<![A-Za-z_$])(\$?[A-Z]{1,3})%d(?![0-9])' % row, r'\1{r}', f)


def fill_col(ws, col, template, r0, r1):
    for r in range(r0, r1 + 1):
        ws[f'{col}{r}'] = template.format(r=r)


def hdr(c, txt, cor=NAVY):
    c.value = txt
    c.fill = fill(cor); c.font = HDR_FONT; c.alignment = HDR_AL; c.border = BOXH


# =====================================================================================
# 1. TABELAS
# =====================================================================================
T = wb['Tabelas']
T['B7'] = datetime.datetime(2026, 1, 1)
T['C7'] = 'A partir desta data valem: alertas, leve atraso e cobertura automática do ferista (padrão EI32: ano todo).'
# parâmetros da macro (linhas 18-19 estavam vazias)
T['A18'] = 'Arquivo da planilha de Pontos'
T['C18'] = 'Caminho completo do arquivo. O botão "Atualizar Pontos" procura "Pontos*.xlsx" na mesma pasta se estiver vazio.'
T['A19'] = 'Última atualização dos pontos'
T['C19'] = 'Preenchido pela macro: data/hora e período da base de pontos lida.'
for a in ('B18', 'B19'):
    T[a].number_format = '@'
nome('cfgArquivoPontos', 'Tabelas!$B$18')
nome('cfgUltimaAtualizacao', 'Tabelas!$B$19')
# bloco novo de parâmetros (padrão EI32), abaixo da tabela de contrato
T['A35'] = 'PARÂMETROS DO PADRÃO EI32'
T['A36'] = 'Doação mínima da sobra do ferista (h)'
T['B36'] = datetime.time(1, 0)
T['C36'] = 'Sobra do ferista abaixo disto NÃO é doada para a próxima falta (igual EI32, Tabelas K15).'
T['A37'] = 'Início do Abatimento de Valor/Função'
T['B37'] = datetime.datetime(2026, 1, 1)
T['C37'] = 'Ferista de cargo menor cobrindo cargo maior: desconta (HH coberto − HH ferista) × horas cobertas a partir desta data.'
T['B36'].number_format = '[h]:mm'
T['B37'].number_format = 'dd/mm/yyyy'
T['A38'] = 'Atualizar Pontos a partir de'
T['B38'] = datetime.datetime(2026, 10, 1)
T['C38'] = ('O botão "Atualizar Pontos" só confere/corrige/cria linhas a partir desta data (protege meses já fechados). '
            'Jan–abr só têm exceções lançadas e mai–jun têm ponto bruto (sem os 30 min de almoço): baixar esta data muda esses meses.')
T['B38'].number_format = 'dd/mm/yyyy'
nome('cfgPontosDesde', 'Tabelas!$B$38')
nome('P_DoacaoMin', 'Tabelas!$B$36')
nome('P_InicioValorFuncao', 'Tabelas!$B$37')

w32_tab = openpyxl.load_workbook(E32)['Tabelas']
# status novos (EI32 v52): modo / conta / trabalhado / categoria / significado
novos = [
    ('Recusa de Atividade S/ Motivo Legítimo', 'Sim', 'Dia inteiro', 'Não', 'Outros',
     'Recusou a atividade sem motivo legítimo. Desconta a jornada. Motivo obrigatório na Observação.'),
    ('Recusa de Atividade C/ Motivo Legítimo', 'Não', 'Dia inteiro', 'Não', 'Justificada',
     'Recusou com motivo legítimo. Não desconta. Motivo obrigatório na Observação.'),
    ('Ausências Devido a Integração / Início de Atividades', 'Não', 'Dia inteiro', 'Não', 'Justificada',
     'Integração TK / início de atividades. Não desconta.'),
    ('Ausências Devido a Afastamento por Acidente', 'Não', 'Dia inteiro', 'Não', 'Justificada',
     'Afastamento por acidente. Não desconta. Motivo obrigatório na Observação.'),
    ('Ausências Devido a Desligamento / Pedido de Demissão', 'Sim', 'Dia inteiro', 'Não', 'Outros',
     'Desligamento ou pedido de demissão: o posto fica descoberto e desconta. Motivo obrigatório.'),
    ('Adicional Não Solicitado', 'Não', 'Não', 'Sim', 'Adicional',
     'Pessoa a mais, não pedida. As horas trabalhadas abatem o absenteísmo do mês 1:1 (como compensação). Não é ferista.'),
]
r = 21
assert T['O20'].value == 'Situação Atipica' and T['O21'].value is None
for st, conta, modo, trab, cat, sig in novos:
    for col, v in zip('OPQRST', (st, conta, modo, trab, cat, sig)):
        c = T[f'{col}{r}']
        c.value = v
        src = T[f'{col}20']
        c.font = src.font.copy(); c.fill = src.fill.copy(); c.border = src.border.copy(); c.alignment = src.alignment.copy()
    r += 1
# tratamentos de HE iguais à EI32
T['Z6'] = 'HORAS COMPENSADAS'
T['Z7'] = 'REMANEJADO P/ OUTRA ÁREA'
for a in ('Z6', 'Z7'):
    T[a].font = T['Z5'].font.copy(); T[a].fill = T['Z5'].fill.copy(); T[a].border = T['Z5'].border.copy()
nome('lstTratHE', 'Tabelas!$Z$4:$Z$7')
# rótulos do relatório (iguais aos da EI32, Tabelas AB:AC) na coluna U, ao lado de "O que significa"
rot32 = {}
for rr in range(3, 41):
    k, v = w32_tab['AB%d' % rr].value, w32_tab['AC%d' % rr].value
    if k and v and k != 'Status':
        rot32[k] = v
rot32['Situação Atipica'] = 'Situação atípica'
T['U3'] = 'RÓTULO NO RELATÓRIO'
T['U3'].font = T['T3'].font.copy(); T['U3'].fill = T['T3'].fill.copy(); T['U3'].alignment = T['T3'].alignment.copy(); T['U3'].border = T['T3'].border.copy()
for rr in range(4, 31):
    st = T['O%d' % rr].value
    if st:
        T['U%d' % rr] = rot32.get(st, st)
        T['U%d' % rr].font = T['T%d' % rr].font.copy(); T['U%d' % rr].border = T['T%d' % rr].border.copy()
T.column_dimensions['U'].width = 34
nome('S_Rotulo', 'Tabelas!$U$4:$U$30')
nome('lstStatus', 'Tabelas!$O$4:$O$30')

# calendário PCAM 2026 (copiado da EI32, Tabelas AE:AG) -> EI36 AF:AH
T32 = w32_tab
T['AF3'], T['AG3'], T['AH3'] = 'CALENDÁRIO PCAM 2026 — DATA', 'DESCRIÇÃO', 'TIPO'
nr = 0
for rr in range(3, 151):
    d = T32[f'AE{rr}'].value
    if not isinstance(d, datetime.datetime):
        continue
    T[f'AF{4 + nr}'] = d; T[f'AF{4 + nr}'].number_format = 'dd/mm/yyyy'
    T[f'AG{4 + nr}'] = T32[f'AF{rr}'].value
    T[f'AH{4 + nr}'] = T32[f'AG{rr}'].value
    nr += 1
nome('tabCalendario', f'Tabelas!$AF$4:$AH${max(4 + nr - 1, 63)}')
T.column_dimensions['AF'].width = 14; T.column_dimensions['AG'].width = 34; T.column_dimensions['AH'].width = 18

# =====================================================================================
# 2. LANCAMENTOS (cabeçalho na linha 3, dados 4..3489)
# =====================================================================================
LW = wb['Lancamentos']
R0 = 4
assert LW['U3'].value.startswith('Jornada') and LW['AU3'].value == 'Compensação pendente'

# --- novas colunas (depois de AU) ---
novas = {
    'AV': ('Prioridade cobertura (aux)', None),
    'AW': ('Carga antes da compensação (aux)', None),
    'AX': ('RG do ferista que cobriu (aux)', None),
    'AY': ('Função do ferista (aux)', None),
    'AZ': ('Valor HH do ferista (aux)', None),
    'BA': ('Abatimento Valor/Função (R$)', None),
    'BB': ('Abatimento Valor/Função — como foi calculado', None),
    'BC': ('Divergência do ponto\n(Atualizar Pontos)', 'in'),
    'BD': ('Conferido?\n(Sim)', 'in'),
    'BE': ('Rel. (aux)', None),
    'BF': ('Rel. nº (aux)', None),
    'BG': ('Valor compensado (R$) (aux)', None),
    'BH': ('Turno 2 automático? (aux)', None),
}
for col, (txt, kind) in novas.items():
    hdr(LW[f'{col}3'], txt, TK if kind == 'in' else NAVY)

F = {}
F['AV'] = '=IF($AQ{r}="","",IF($C{r}="Ausências Devido a Férias",1000,0)+MAX($W{r}-N($M{r}),0)*24)'
F['AW'] = '=IF($A{r}="",0,MAX($Y{r}-$Z{r}-$AB{r},0))'
F['AX'] = ('=IF(N($Z{r})<=0,"",IF(COUNTIFS(L_Data,$A{r},L_Equipe,$P{r},L_Contrato,"Ferista")=1,'
           'SUMIFS(L_RG,L_Data,$A{r},L_Equipe,$P{r},L_Contrato,"Ferista"),'
           'IFERROR(LOOKUP(2,1/((V_Equipe=$P{r})*(V_Contrato="Ferista")*(V_Ini<=$A{r})),V_RG),"")))')
F['AY'] = '=IF($AX{r}="","",IFERROR(LOOKUP(2,1/((V_RG=$AX{r})*(V_Ini<=$A{r})),V_Funcao),""))'
F['AZ'] = ('=IF($AY{r}="","",IFERROR(INDEX(F_Valor,MATCH($AY{r}&"|"&SUMPRODUCT(MAX((F_Funcao=$AY{r})*(F_Ini<=$A{r})*F_Ini)),'
           'F_Chave,0)),""))')
F['BA'] = ('=IF(OR($A{r}="",$AZ{r}="",N($Z{r})<=0,NOT(ISNUMBER($T{r})),$Q{r}="Ferista",ISNUMBER(SEARCH("ENCARREGADO",$R{r}))),0,'
           'IF($A{r}<P_InicioValorFuncao,0,ROUND(MAX(ROUND($T{r},2)-ROUND($AZ{r},2),0)*N($Z{r})*24,2)))')
F['BB'] = ('=IF(N($BA{r})<=0,"","Acobertado pelo ferista / Abatimento de Valor/Função: ferista "&$AY{r}&" (R$ "&FIXED($AZ{r},2)'
           '&"/h) cobrindo "&$R{r}&" (R$ "&FIXED($T{r},2)&"/h). Diferença R$ "&FIXED(ROUND($T{r},2)-ROUND($AZ{r},2),2)'
           '&"/h × "&TEXT($Z{r},"[h]:mm")&"h cobertas = R$ "&FIXED($BA{r},2)&". Não gera absenteísmo.")')
F['BE'] = ('=IF(AND($A{r}<>"",$N{r}=Painel!$C$3,$P{r}=P_EquipePrincipal,OR($C{r}<>"Horas Trabalhadas",N($Y{r})>0,N($Z{r})>0,'
           'N($AE{r})>0,$J{r}<>"",N($M{r})>0,N($L{r})>0,AND($BC{r}<>"",$BD{r}<>"Sim"))),1,0)')
F['BF'] = '=IF($BE{r}=1,COUNTIF($BE$4:$BE{r},1),0)'
F['BG'] = '=IF(OR(N($AA{r})=0,$T{r}=""),0,$AA{r}*24*$T{r})'
F['BH'] = ('=IF(AND($A{r}<>"",$I{r}="",UPPER(TRIM($S{r}))="ADM",WEEKDAY($A{r},2)<=5,ISNUMBER(SEARCH(" = ",$D{r})),'
           'IFERROR(TIMEVALUE(LEFT($D{r},5)),1)<TIME(7,0,0),'
           'IFERROR(TIMEVALUE(RIGHT(LEFT($D{r},SEARCH(" = ",$D{r})-1),5)),1)<=TIME(14,0,0)),"Sim","")')

# --- colunas existentes que mudam ---
# U: jornada com Turno 2 automático (EI32 AD)
F['U'] = '=IF(OR($A{r}="",$S{r}=""),0,IFERROR(INDEX(T_Horas,MATCH(IF($BH{r}="Sim","TURNO 2",$S{r}),T_Turno,0),WEEKDAY($A{r},2)),0))'
# AR: ordem pela prioridade EI32 (férias primeiro, depois a maior falta)
F['AR'] = '=IF($AQ{r}="","",COUNTIFS(L_ChaveCob,$AQ{r},L_Prior,">"&$AV{r})+COUNTIFS($AQ$4:$AQ{r},$AQ{r},$AV$4:$AV{r},$AV{r}))'
# AT: cascata com doação mínima (EI32 AC)
cap = '($AS{r}-SUMIFS(L_Ausencia,L_ChaveCob,$AQ{r},L_OrdemCob,"<"&$AR{r}))'
F['AT'] = (f'=IF(OR($AQ{{r}}="",$AS{{r}}=0),0,IF({cap}<P_DoacaoMin,0,IF({cap}>=$W{{r}}-P_DoacaoMin,$W{{r}},MAX({cap},0))))')
# AB antes de AA (acordo antes da compensação, como na EI32: AM = V − AC + M − AJ − Q, compensação no mês)
F['AB'] = '=IF($A{r}="",0,MIN(MAX($Y{r}-$Z{r},0),N($M{r})))'
# AA: compensação em bloco — fatia proporcional do mês (Resumo_Mensal AN)
F['AA'] = ('=IF(OR($A{r}="",N($AW{r})=0,$P{r}<>P_EquipePrincipal),0,'
           '$AW{r}*IFERROR(INDEX(Resumo_Mensal!$AN$5:$AN$16,MATCH($N{r},Resumo_Mensal!$A$5:$A$16,0)),0))')
F['AC'] = '=IF($A{r}="",0,MAX($AW{r}-$AA{r},0))'

# AG (de onde vem) e AH (alerta): partem da fórmula do Felipe
ag = tpl(LW['AG4'].value)
old_comp = ('IF($AA{r}>0,"Compensado "&TEXT($AA{r},"[h]:mm")&"h por "&IFERROR(INDEX(CP_Quem,MATCH($AP{r},CP_ChaveFalta,0)),"?")&". ","")')
assert old_comp in ag, 'AG mudou'
ag = ag.replace(old_comp, 'IF($AA{r}>0,"Compensação do mês (em bloco, proporcional): "&TEXT($AA{r},"[h]:mm")&"h. ","")')
old_fer = 'IF($Z{r}>0,"Ferista cobriu "&TEXT($Z{r},"[h]:mm")&"h. ","")'
assert old_fer in ag
ag = ag.replace(old_fer, 'IF($Z{r}>0,IF($Z{r}<$W{r},"Sobra do ferista doada: "&TEXT($Z{r},"[h]:mm")&"h de "&TEXT($W{r},"[h]:mm")&"h. ",'
                         '"Ferista cobriu "&TEXT($Z{r},"[h]:mm")&"h. "),"")&IF($BH{r}="Sim","Turno 2 automático (entrou antes das 07:00 e saiu até 14:00). ","")')
old_desc = 'IF($AC{r}>0,"DESCONTA "'
assert old_desc in ag
ag = ag.replace(old_desc, 'IF($BA{r}>0,"Valor/Função: R$ "&FIXED($BA{r},2)&". ","")&IF($AC{r}>0,"DESCONTA "')
F['AG'] = ag

ah = tpl(LW['AH4'].value)
old_vinc = '"VINCULAR: registrar na aba Compensações de quem é a falta"'
assert old_vinc in ah
ah = ah.replace(old_vinc, '"VINCULAR: registrar na aba Compensações (quem trabalhou, horas e aprovação)"')
old_tail = 'IF(AND($J{r}<>"",N($K{r})+N($L{r})>$U{r}),"CONFERIR: horas manuais maiores que a jornada","")'
assert old_tail in ah
turno_alto = ('AND($I{r}="",UPPER(TRIM($S{r}))="ADM",WEEKDAY($A{r},2)<=5,ISNUMBER(SEARCH(" = ",$D{r})),'
              'IFERROR(TIMEVALUE(LEFT($D{r},5)),1)<TIME(7,0,0),'
              'IFERROR(TIMEVALUE(RIGHT(LEFT($D{r},SEARCH(" = ",$D{r})-1),5)),1)>TIME(14,0,0))')
ah = ah.replace(old_tail, 'IF(AND($J{r}<>"",N($K{r})+N($L{r})>$U{r}),"CONFERIR: horas manuais maiores que a jornada",'
                          'IF(AND($BC{r}<>"",$BD{r}<>"Sim"),"PONTO DIVERGENTE: conferir e marcar Conferido? = Sim (col. BD)",'
                          f'IF({turno_alto},"TURNO: entrou antes das 07:00 e saiu depois das 14:00 — pode ser compensação, conferir","")))')
F['AH'] = ah

for col, t in F.items():
    fill_col(LW, col, t, R0, LAST)

# estilo das colunas novas (copia do corpo da AU / AG)
for col in novas:
    src = 'AG' if col in ('BB', 'BC') else 'AU'
    for rr in range(R0, LAST + 1):
        c, s = LW[f'{col}{rr}'], LW[f'{src}{rr}']
        c.font = s.font.copy(); c.border = s.border.copy(); c.alignment = s.alignment.copy(); c.number_format = s.number_format
for col in ('AV', 'AW'):
    for rr in range(R0, LAST + 1):
        LW[f'{col}{rr}'].number_format = '[h]:mm' if col == 'AW' else '0.00'
for col in ('AZ', 'BA', 'BG'):
    for rr in range(R0, LAST + 1):
        LW[f'{col}{rr}'].number_format = '#,##0.00'
for col in ('BC', 'BD'):
    for rr in range(R0, LAST + 1):
        LW[f'{col}{rr}'].number_format = 'General'
LW.column_dimensions['BA'].width = 14
LW.column_dimensions['BB'].width = 60
LW.column_dimensions['BC'].width = 60
LW.column_dimensions['BD'].width = 11
for col in ('AV', 'AW', 'AX', 'AY', 'AZ', 'BE', 'BF', 'BG', 'BH'):
    LW.column_dimensions[col].width = 12
    LW.column_dimensions[col].outlineLevel = 1
    LW.column_dimensions[col].hidden = True
dv = DataValidation(type='list', formula1='lstSimNao', allow_blank=True)
dv.add(f'BD4:BD{LAST}')
LW.add_data_validation(dv)
LW.conditional_formatting.add(f'BC4:BD{LAST}', FormulaRule(formula=['AND($BC4<>"",$BD4<>"Sim")'], fill=fill(VERM),
                                                            font=Font(color='FFC00000', bold=True)))
LW.conditional_formatting.add(f'BC4:BD{LAST}', FormulaRule(formula=['AND($BC4<>"",$BD4="Sim")'], fill=fill(VERDE)))
LW.conditional_formatting.add(f'BA4:BB{LAST}', FormulaRule(formula=['N($BA4)>0'], fill=fill(VERM)))
LW.auto_filter.ref = f'A3:BD{LAST}'
LW['AJ2'] = 'AUXILIARES (não editar) — clique no [+] acima para ver'
LW['A2'] = ('CÉLULA CINZA-CLARO / CABEÇALHO AZUL = você preenche (A–G no dia a dia; H–M só em casos especiais).  '
            'CABEÇALHO MARINHO = automático, não digite por cima.  AG explica de onde veio o desconto; AH mostra o que precisa de atenção; '
            'BC/BD = conferência do ponto.')

# nomes novos de Lancamentos
for n, c in {'L_Prior': 'AV', 'L_CargaPre': 'AW', 'L_ValorFuncao': 'BA', 'L_Div': 'BC', 'L_Conf': 'BD',
             'L_Rel': 'BE', 'L_RelN': 'BF', 'L_ValorComp': 'BG', 'L_Explica': 'AG', 'L_Calculo': 'BB',
             'L_Turno': 'S', 'L_Valorhh': 'T', 'L_Ponto': 'D', 'L_Obs': 'G'}.items():
    nome(n, f'Lancamentos!${c}$4:${c}${LAST}')

# =====================================================================================
# 3. COMPENSACOES (em bloco, como a EI32)
# =====================================================================================
CW = wb['Compensacoes']
assert CW['J3'].value.startswith('Total aprovado') and CW['K3'].value.startswith('Abatido')
CW['A2'] = ('Cada linha: QUEM trabalhou para compensar, QUANDO e QUANTAS horas. Só conta com "Aprovado? = Sim". '
            'As horas aprovadas do mês abatem o absenteísmo da equipe em bloco (1:1, igual à EI32). '
            '"Compensa a falta de" e "Data da falta" são opcionais (só para rastreio).')
CW['D3'] = 'Compensa a falta de\n(opcional)'
CW['E3'] = 'Data da falta\n(opcional)'
CW['I3'] = 'Falta informada (h)'
CW['J3'] = 'Horas que contam\n(aprovadas)'
CW['K3'] = 'Mês'
hdr(CW['O3'], 'Equipe (aux)', NAVY)
for c in ('I3', 'J3', 'K3'):
    hdr(CW[c], CW[c].value, NAVY)
for rr in range(4, CLAST + 1):
    CW[f'J{rr}'] = f'=IF(AND($F{rr}="Sim",$A{rr}<>"",TRIM($B{rr})<>""),N($C{rr}),0)'
    CW[f'K{rr}'] = f'=IF($A{rr}="","",DATE(YEAR($A{rr}),MONTH($A{rr}),1))'
    CW[f'O{rr}'] = f'=IF(TRIM($B{rr})="","",IFERROR(INDEX(Cadastro!$E$4:$E$200,MATCH(TRIM($B{rr}),C_Nome,0)),"?"))'
    CW[f'L{rr}'] = (f'=IF($A{rr}&$B{rr}&$C{rr}&$D{rr}&$E{rr}="","",IF(OR($A{rr}="",TRIM($B{rr})="",$C{rr}=""),"PREENCHER: data, quem trabalhou e horas",'
                    f'IF($O{rr}="?","CONFERIR: funcionário fora do Cadastro",IF($F{rr}="","PREENCHER: aprovado? (Sim / Não / Pendente)",'
                    f'IF($F{rr}="Pendente","PENDENTE: aguardando aprovação (ainda não abate)",IF($F{rr}<>"Sim","NÃO APROVADA: não abate",'
                    f'IF(AND($N{rr}<>"",COUNTIF(L_ChavePD,$N{rr})=0),"CONFERIR: a falta informada não existe nos Lançamentos",'
                    f'"OK: abate o absenteísmo do mês (em bloco)")))))))')
    CW[f'P{rr}'] = f'=IF(AND($A{rr}<>"",$K{rr}=Painel!$C$3),1,0)'
    CW[f'Q{rr}'] = f'=IF($P{rr}=1,COUNTIF($P$4:$P{rr},1),0)'
    for col in 'JKOPQ':
        s = CW[f'I{rr}']
        c = CW[f'{col}{rr}']
        c.font = s.font.copy(); c.border = s.border.copy(); c.alignment = s.alignment.copy()
    CW[f'J{rr}'].number_format = '[h]:mm'
    CW[f'K{rr}'].number_format = 'mmm/yy'
hdr(CW['P3'], 'Rel. (aux)', NAVY); hdr(CW['Q3'], 'Rel. nº (aux)', NAVY)
for col in ('M', 'N', 'O', 'P', 'Q'):
    CW.column_dimensions[col].outlineLevel = 1
    CW.column_dimensions[col].hidden = True
CW.column_dimensions['L'].width = 48
nome('CP_Conta', f'Compensacoes!$J$4:$J${CLAST}')
nome('CP_Mes', f'Compensacoes!$K$4:$K${CLAST}')
nome('CP_Equipe', f'Compensacoes!$O$4:$O${CLAST}')
nome('CP_Obs', f'Compensacoes!$G$4:$G${CLAST}')
nome('CP_RelN', f'Compensacoes!$Q$4:$Q${CLAST}')
nome('CP_Tipo', f'Compensacoes!$H$4:$H${CLAST}')
dva = DataValidation(type='list', formula1='lstAprov', allow_blank=True)
dva.add(f'F4:F{CLAST}')
CW.add_data_validation(dva)

# =====================================================================================
# 4. LANC_HE: divergência / conferido / registrar compensação / relatório
# =====================================================================================
HW = wb['Lanc_HE']
assert HW['AB3'].value == 'Alerta' and HW['AD3'].value.startswith('vig#')
hdr(HW['AE3'], 'Divergência do ponto\n(Atualizar Pontos)', TK)
hdr(HW['AF3'], 'Conferido?\n(Sim)', TK)
hdr(HW['AG3'], 'Rel. HE (aux)', NAVY)
hdr(HW['AH3'], 'Rel. HE nº (aux)', NAVY)
for rr in range(4, HLAST + 1):
    HW[f'AB{rr}'] = (f'=IF($A{rr}="","",IF($O{rr}="?","CADASTRO: funcionário não está no Cadastro",IF($G{rr}="","DEFINIR: tratamento da HE (EI, área, compensadas ou remanejado)",'
                     f'IF(AND($G{rr}="HORAS COMPENSADAS",COUNTIF(CP_ChaveQuem,TRIM($B{rr})&"|"&$A{rr})=0),"REGISTRAR: lançar na aba Compensações (quem compensou e se foi aprovado)",'
                     f'IF(AND($AE{rr}<>"",$AF{rr}<>"Sim"),"PONTO DIVERGENTE: conferir e marcar Conferido? = Sim (col. AF)",'
                     f'IF($L{rr}<>"Sim","PONTO: aguardando histórico dos pontos",""))))))')
    HW[f'AG{rr}'] = f'=IF(AND($A{rr}<>"",$N{rr}=Painel!$C$3),1,0)'
    HW[f'AH{rr}'] = f'=IF($AG{rr}=1,COUNTIF($AG$4:$AG{rr},1),0)'
    for col in ('AE', 'AF', 'AG', 'AH'):
        s = HW[f'AB{rr}']; c = HW[f'{col}{rr}']
        c.font = s.font.copy(); c.border = s.border.copy(); c.alignment = s.alignment.copy()
HW.column_dimensions['AE'].width = 60; HW.column_dimensions['AF'].width = 11
for col in ('AG', 'AH'):
    HW.column_dimensions[col].outlineLevel = 1; HW.column_dimensions[col].hidden = True
dvh = DataValidation(type='list', formula1='lstTratHE', allow_blank=True); dvh.add(f'G4:G{HLAST}'); HW.add_data_validation(dvh)
dvc = DataValidation(type='list', formula1='lstSimNao', allow_blank=True); dvc.add(f'AF4:AF{HLAST}'); HW.add_data_validation(dvc)
HW.conditional_formatting.add(f'AE4:AF{HLAST}', FormulaRule(formula=['AND($AE4<>"",$AF4<>"Sim")'], fill=fill(VERM), font=Font(color='FFC00000', bold=True)))
HW.conditional_formatting.add(f'AE4:AF{HLAST}', FormulaRule(formula=['AND($AE4<>"",$AF4="Sim")'], fill=fill(VERDE)))
for n, c in {'H_Div': 'AE', 'H_Conf': 'AF', 'H_RelN': 'AH', 'H_Data': 'A', 'H_Trat': 'G', 'H_Tipo': 'T', 'H_Ponto': 'K'}.items():
    nome(n, f'Lanc_HE!${c}$4:${c}${HLAST}')

# =====================================================================================
# 5. RESUMO_MENSAL: compensação em bloco + Valor/Função
# =====================================================================================
RW = wb['Resumo_Mensal']
assert RW['AH4'].value.startswith('Descontado: Situação') and RW['AI4'].value is None
newR = {
    'AI': ('Absenteísmo antes da compensação (h)', '=SUMIFS(L_CargaPre,L_Mes,$A{r},L_Equipe,P_EquipePrincipal)', '[h]:mm'),
    'AJ': ('Compensação aprovada no mês (h)', '=SUMIFS(CP_Conta,CP_Mes,$A{r},CP_Equipe,P_EquipePrincipal)', '[h]:mm'),
    'AK': ('Adicional não solicitado (h)', '=SUMIFS(L_Trab,L_Mes,$A{r},L_Equipe,P_EquipePrincipal,L_Status,"Adicional Não Solicitado")', '[h]:mm'),
    'AL': ('Compensação usada (h)', '=MIN(N($AJ{r}),N($AI{r}))', '[h]:mm'),
    'AM': ('Adicional usado (h)', '=MIN(N($AK{r}),MAX(N($AI{r})-N($AL{r}),0))', '[h]:mm'),
    'AN': ('Fator compensado (aux)', '=IF(N($AI{r})<=0,0,(N($AL{r})+N($AM{r}))/N($AI{r}))', '0.0000'),
    'AO': ('Sobra de compensação — horas positivas (h)', '=MAX(N($AJ{r})-N($AL{r}),0)', '[h]:mm'),
    'AP': ('Valor compensado (R$)', '=SUMIFS(L_ValorComp,L_Mes,$A{r},L_Equipe,P_EquipePrincipal)', '#,##0.00'),
    'AQ': ('Abatimento Valor/Função (R$)', '=SUMIFS(L_ValorFuncao,L_Mes,$A{r},L_Equipe,P_EquipePrincipal)', '#,##0.00'),
}
for col, (h, f, fmt) in newR.items():
    c = RW[f'{col}4']; c.value = h
    s = RW['AH4']; c.font = s.font.copy(); c.fill = s.fill.copy(); c.border = s.border.copy(); c.alignment = s.alignment.copy()
    for rr in range(5, 17):
        x = RW[f'{col}{rr}']; x.value = f.format(r=rr)
        s = RW[f'AH{rr}']; x.font = s.font.copy(); x.border = s.border.copy(); x.alignment = s.alignment.copy(); x.number_format = fmt
    t = RW[f'{col}17']; s = RW['AH17']
    t.value = '' if col == 'AN' else f'=SUM({col}5:{col}16)'
    t.font = s.font.copy(); t.fill = s.fill.copy(); t.border = s.border.copy(); t.number_format = fmt
    RW.column_dimensions[col].width = 16
for rr in range(5, 17):
    assert RW[f'H{rr}'].value == f'=SUMIFS(L_Valor,L_Mes,$A{rr},L_Equipe,P_EquipePrincipal)'
    RW[f'H{rr}'] = f'=SUMIFS(L_Valor,L_Mes,$A{rr},L_Equipe,P_EquipePrincipal)+N($AQ{rr})'
    RW[f'M{rr}'] = f'=SUMIFS(CP_Horas,CP_Mes,$A{rr},CP_Equipe,P_EquipePrincipal,CP_Aprov,"<>Sim")'
RW['E4'] = '(−) Compensado + adicional (h)'
RW['H4'] = 'Abatimento (R$)\n(inclui Valor/Função)'
RW['M4'] = 'Compensação não aprovada (h)'
RW['A2'] = ('Tudo aqui é automático, exceto a coluna "Desconto SLA". Leia da esquerda para a direita: brutas − ferista − perdoado − '
            'compensado (em bloco, proporcional) = descontadas; Abatimento = valor descontado + Abatimento de Valor/Função (col. AQ).')
for n, c in {'R_Mes': 'A', 'R_Contrato': 'B', 'R_Bruto': 'C', 'R_Ferista': 'D', 'R_Comp': 'E', 'R_Perdoado': 'F', 'R_Desc': 'G',
             'R_Abat': 'H', 'R_SLA': 'I', 'R_Final': 'J', 'R_Just': 'K', 'R_ValorJust': 'L', 'R_CompPend': 'M', 'R_FerTrab': 'V',
             'R_FerSem': 'W', 'R_CompAprov': 'AJ', 'R_Adic': 'AK', 'R_CompUsada': 'AL', 'R_AdicUsado': 'AM', 'R_Sobra': 'AO',
             'R_ValorComp': 'AP', 'R_ValorFuncao': 'AQ', 'R_HEEI': 'R', 'R_Qtd': 'Y'}.items():
    nome(n, f'Resumo_Mensal!${c}$5:${c}$16')

# =====================================================================================
# 6. PAINEL
# =====================================================================================
P = wb['Painel']
assert P['B20'].value == 'TOTAL ABATIDO' and P['B21'].value.startswith('Confere')
idx = lambda col: f'IFERROR(INDEX(Resumo_Mensal!${col}$5:${col}$16,MATCH($C$3,Resumo_Mensal!$A$5:$A$16,0)),0)'
# linha 20 -> Valor/Função ; 21 -> TOTAL ; 22 -> confere
for col in 'BCDEFGHI':
    s20, d21 = P[f'{col}20'], P[f'{col}21']
    d21.font = s20.font.copy(); d21.fill = s20.fill.copy(); d21.border = s20.border.copy(); d21.alignment = s20.alignment.copy(); d21.number_format = s20.number_format
    s19 = P[f'{col}19']
    s20.font = s19.font.copy(); s20.fill = s19.fill.copy(); s20.border = s19.border.copy(); s20.alignment = s19.alignment.copy(); s20.number_format = s19.number_format
    P[f'{col}22'].font = Font(name='Arial', size=9, bold=True, color=TXT)
P['B20'] = '( + ) Abatimento de Valor/Função (ferista de cargo menor)'
for col in 'CDEFG':
    P[f'{col}20'] = None
P['H20'] = '=' + idx('AQ')
P['I20'] = '=IF($H$21=0,0,H20/$H$21)'
P['B21'] = 'TOTAL ABATIDO'
for col in 'CDEFG':
    P[f'{col}21'] = f'=SUM({col}10:{col}19)'
P['H21'] = '=SUM(H10:H20)'
P['I21'] = '=IF($H$21=0,0,H21/$H$21)'
P['B22'] = 'Confere com o Resumo_Mensal?'
P['H22'] = '=IF(ROUND(H21-D6,2)=0,"✔ confere","✘ diferença")'
P['H22'].alignment = Alignment(horizontal='center')
for rr in range(10, 20):
    P[f'I{rr}'] = f'=IF($H$21=0,0,H{rr}/$H$21)'
P['B8'] = 'DE ONDE VEM O ABATIMENTO  (horas lançadas − ferista − compensação (em bloco) − acordo = descontado)'
# compensações (em bloco)
P['F24'] = 'Compensação + adicional usados no mês (h)'
P['I24'] = '=' + idx('E')
P['F25'] = 'Compensação não aprovada no mês (h)'
P['I25'] = '=' + idx('M')
P['F28'] = 'Sobra de compensação — horas positivas (h)'
P['I28'] = '=' + idx('AO')
P.merge_cells('F28:H28')
for a, s in (('F28', 'F27'), ('I28', 'I27')):
    P[a].font = P[s].font.copy(); P[a].border = P[s].border.copy(); P[a].alignment = P[s].alignment.copy(); P[a].number_format = P['I24'].number_format
# ferista: valor/função
P['B34'] = 'Abatimento de Valor/Função (R$)'
P['C34'] = '=' + idx('AQ')
for a, s in (('B34', 'B33'), ('C34', 'C33')):
    P[a].font = P[s].font.copy(); P[a].border = P[s].border.copy(); P[a].alignment = P[s].alignment.copy()
P['C34'].number_format = '#,##0.00'
# verificação: divergências de ponto
P['K19'] = 'Ponto divergente ainda não conferido (Atualizar Pontos)'
P['L19'] = '=COUNTIFS(L_Div,"?*",L_Mes,$C$3,L_Conf,"<>Sim")+COUNTIFS(H_Div,"?*",H_Mes,$C$3,H_Conf,"<>Sim")'
for a, s in (('K19', 'K17'), ('L19', 'L17')):
    P[a].font = P[s].font.copy(); P[a].border = P[s].border.copy(); P[a].alignment = P[s].alignment.copy(); P[a].fill = P[s].fill.copy()
P['K18'] = 'Filtre a coluna AH de Lançamentos pelo texto do alerta para achar as linhas.'
P['L18'] = None
# menu (linha 2) e espaço dos botões (linha 3)
menu = [('B2', 'Relatorio', 'Relatório do mês (PDF) ►'), ('C2', 'Lancamentos', 'Lançamentos ►'),
        ('D2', 'Compensacoes', 'Compensações ►'), ('F2', 'Lanc_HE', 'Horas extras ►'),
        ('H2', 'Conferencia_Ponto', 'Conferência de ponto ►')]
for a, aba, txt in menu:
    P[a] = f'=HYPERLINK("#\'{aba}\'!A1","{txt}")'
for a in ('B2', 'C2', 'D2', 'F2', 'H2'):
    P[a].fill = fill(TK); P[a].font = Font(name='Arial', size=9, bold=True, color='FFFFFFFF')
    P[a].alignment = Alignment(horizontal='center', vertical='center'); P[a].border = BOXH
for a in ('E2', 'G2', 'I2'):
    P[a].fill = fill(TK); P[a].border = BOXH
P.merge_cells('D2:E2'); P.merge_cells('F2:G2'); P.merge_cells('H2:I2')
P.row_dimensions[2].height = 20
P.row_dimensions[3].height = 26

# =====================================================================================
# 7. CADASTRO: matrícula real no lugar do RG provisório 9000xx
# =====================================================================================
def sem_acento(s):
    return ''.join(c for c in unicodedata.normalize('NFD', str(s).upper().strip()) if unicodedata.category(c) != 'Mn')

wp = openpyxl.load_workbook(PONTOS, read_only=True)
mat = {}
for row in wp['Base unificada'].iter_rows(min_row=5, values_only=True):
    if row[0]:
        mat.setdefault(sem_acento(row[0]), set()).add(str(row[1]))
CAD = wb['Cadastro']
trocas = []
for rr in range(4, 201):
    rg, nm = CAD[f'A{rr}'].value, CAD[f'B{rr}'].value
    if isinstance(rg, int) and 900000 <= rg < 901000 and nm:
        m = mat.get(sem_acento(nm))
        if m and len(m) == 1:
            novo = int(next(iter(m)))
            CAD[f'A{rr}'] = novo
            obs = CAD[f'D{rr}'].value
            CAD[f'D{rr}'] = ((obs + ' | ') if obs else '') + f'RG provisório {rg} trocado pela matrícula do relógio ({novo}) em 23/10/2026.'
            trocas.append((nm, rg, novo))
print('RGs trocados:', trocas)

# =====================================================================================
# 8. ABA CONFERENCIA_PONTO (preenchida pela macro)
# =====================================================================================
CF = wb.create_sheet('Conferencia_Ponto', index=wb.sheetnames.index('Compensacoes') + 1)
CF.sheet_view.showGridLines = False
for col, w in {'A': 1.6, 'B': 11, 'C': 13, 'D': 32, 'E': 12, 'F': 30, 'G': 26, 'H': 26, 'I': 28, 'J': 12, 'K': 1.6, 'L': 40}.items():
    CF.column_dimensions[col].width = w
CF['B1'] = '=HYPERLINK("#\'Painel\'!A1","CONFERÊNCIA DE PONTO   ◄ Painel")'
CF['B1'].font = Font(name='Arial', size=16, bold=True, color=TK)
CF.row_dimensions[1].height = 30
CF.merge_cells('B1:F1')   # G1:J1 = botões da macro
CF['L1'] = 'Clique na data para ir até a linha.'
CF['L1'].font = Font(name='Arial', size=8, italic=True, color=CINZA)
CF['B2'] = '=IF(cfgUltimaAtualizacao="","Ainda não atualizado.","Última atualização: "&cfgUltimaAtualizacao)'
CF['B2'].font = Font(name='Arial', size=10, bold=True, color=TXT)
CF.merge_cells('B2:J2')
CF['B3'] = ('Lista gerada pelo botão "Atualizar Pontos" (Suplementos › Controle de Horas TK). Cada linha vermelha em Lançamentos (col. BC) '
            'e em Lanc_HE (col. AE) aparece aqui. Analise e marque "Conferido? = Sim" na própria linha.')
CF['B3'].font = Font(name='Arial', size=8, color=CINZA)
CF['B3'].alignment = Alignment(wrap_text=True, vertical='top')
CF.merge_cells('B3:J3'); CF.row_dimensions[3].height = 26
for col, txt in zip('BCDEFIJ', ('Data', 'Aba', 'Funcionário', 'Equipe', 'Divergência (texto da coluna de divergência)', 'Situação', 'Conferido?')):
    hdr(CF[f'{col}5'], txt, TK)
for col in 'GH':
    CF[f'{col}5'].fill = fill(TK); CF[f'{col}5'].border = BOXH
CF.merge_cells('F5:H5')
CF.freeze_panes = 'B6'
CF.sheet_properties.tabColor = NAVY

# =====================================================================================
# 9. ABA RELATORIO (fonte do PDF da macro CriarRelatorio)
# =====================================================================================
RL = wb.create_sheet('Relatorio', index=wb.sheetnames.index('Painel') + 1)
RL.sheet_view.showGridLines = False
for col, w in {'A': 1.6, 'B': 11.5, 'C': 30, 'D': 24, 'E': 9.5, 'F': 9.5, 'G': 9.5, 'H': 9.5, 'I': 52, 'J': 13, 'K': 1.6,
               'L': 7.7, 'M': 4, 'N': 3}.items():
    RL.column_dimensions[col].width = w
M3 = 'Painel!$C$3'
ix = lambda col: f'IFERROR(INDEX(Resumo_Mensal!${col}$5:${col}$16,MATCH({M3},Resumo_Mensal!$A$5:$A$16,0)),0)'
t = lambda sz=9, b=False, cor=TXT, it=False: Font(name='Arial', size=sz, bold=b, color=cor, italic=it)


def faixa(row, txt):
    RL[f'B{row}'] = txt
    for col in 'BCDEFGHIJ':
        RL[f'{col}{row}'].fill = fill(TK); RL[f'{col}{row}'].border = BOXH
    RL[f'B{row}'].font = Font(name='Arial', size=10, bold=True, color='FFFFFFFF')
    RL[f'B{row}'].alignment = Alignment(horizontal='center', vertical='center')
    RL.merge_cells(f'B{row}:J{row}')
    RL.row_dimensions[row].height = 18


def cab(row, cols):
    for col, txt in cols.items():
        c = RL[f'{col}{row}']; c.value = txt
        c.fill = fill(F2); c.font = t(9, True); c.alignment = HDR_AL; c.border = BOX
    RL.row_dimensions[row].height = 26


def cel(a, v, fmt=None, b=False, al='center', sz=9):
    c = RL[a]; c.value = v; c.font = t(sz, b); c.border = BOX
    c.alignment = Alignment(horizontal=al, vertical='center', wrap_text=(al == 'left'))
    if fmt: c.number_format = fmt
    return c


RL['B1'] = '=HYPERLINK("#\'Painel\'!A1","EI36 SERRALHERIA — RELATÓRIO DE OCORRÊNCIAS DO MÊS   ◄ Painel")'
RL['B1'].font = Font(name='Arial', size=16, bold=True, color=TK); RL['B1'].alignment = Alignment(horizontal='left', vertical='center')
RL.merge_cells('B1:H1'); RL.row_dimensions[1].height = 30   # I1:J1 = botão "Criar Relatório (PDF)"
RL['B2'] = (f'="Mês: "&CHOOSE(MONTH({M3}),"janeiro","fevereiro","março","abril","maio","junho","julho","agosto","setembro",'
            f'"outubro","novembro","dezembro")&"/"&YEAR({M3})&"   ·   emitido em "&TEXT(TODAY(),"dd/mm/yyyy")')
RL['B2'].font = t(10, True); RL['B2'].alignment = Alignment(horizontal='center'); RL.merge_cells('B2:J2')
RL['B3'] = ('Tudo o que não foi um dia normal de trabalho: faltas, atrasos (leves ou não), férias, atestados, remanejamentos, '
            'cobertura do ferista, compensações e horas extras. Escolha o mês no Painel (C3).')
RL['B3'].font = t(8, cor=CINZA); RL['B3'].alignment = Alignment(horizontal='center', wrap_text=True)
RL.merge_cells('B3:J3'); RL.row_dimensions[3].height = 24
RL['L2'] = f'=YEAR({M3})&"-"&IF(MONTH({M3})<10,"0","")&MONTH({M3})'

# --- resumo ---
faixa(5, 'RESUMO DO MÊS')
cab(6, {'B': 'Equipe', 'C': 'Contrato (R$)', 'D': 'Abatimento (R$)', 'E': 'Ocorr.', 'F': 'Horas desc.', 'G': 'Ferista cobriu',
        'H': 'Faltas just.', 'I': 'Valor final a pagar (R$)', 'J': 'HE pagas por EI (R$)'})
cel('B7', '=P_EquipePrincipal', b=True)
cel('C7', '=' + ix('B'), '#,##0.00')
cel('D7', '=' + ix('H'), '#,##0.00', b=True)
cel('E7', '=SUM(L_Rel)', '0')
cel('F7', '=' + ix('G'), '[h]:mm')
cel('G7', '=' + ix('D'), '[h]:mm')
cel('H7', '=' + ix('K'), '[h]:mm')
cel('I7', '=' + ix('J'), '#,##0.00', b=True)
cel('J7', '=' + ix('R'), '#,##0.00')
RL['B8'] = f'="Desconto SLA do mês (já descontado do valor final): R$ "&FIXED({ix("I")},2)'
RL['B8'].font = t(9, True, 'FFC00000'); RL.merge_cells('B8:J8')
RL['N8'] = f'=IF({ix("I")}>0,"",1)'
RL['B9'] = (f'="Composição: (−) ausências R$ "&FIXED({ix("H")}-{ix("AQ")},2)&"   (+) Abatimento de Valor/Função R$ "&FIXED({ix("AQ")},2)'
            f'&"   (−) SLA R$ "&FIXED({ix("I")},2)&"   =   valor final R$ "&FIXED({ix("J")},2)')
RL['B9'].font = t(8, cor=CINZA); RL.merge_cells('B9:J9')

# --- pendências para fechar ---
faixa(11, 'PENDÊNCIAS PARA FECHAR O MÊS (tudo precisa estar em zero)')
cab(12, {'B': 'Qtde', 'C': 'Pendência', 'I': 'O que fazer'})
RL.merge_cells('C12:H12'); RL.merge_cells('I12:J12')
pend = [
    (f'=COUNTIFS(L_Div,"?*",L_Mes,{M3},L_Conf,"<>Sim")', 'Lançamentos: ponto divergente ainda não conferido', 'Analisar a linha vermelha e marcar "Conferido? = Sim" (coluna BD).'),
    (f'=COUNTIFS(H_Div,"?*",H_Mes,{M3},H_Conf,"<>Sim")', 'Lanc_HE: horas extras com ponto divergente não conferidas', 'Conferir o ponto e marcar "Conferido? = Sim" (coluna AF).'),
    (f'=COUNTIFS(H_Mes,{M3},H_Trat,"")', 'Horas extras sem tratamento definido', 'Escolher o tratamento (pagas por EI, pela área, compensadas ou remanejadas).'),
    (f'=COUNTIFS(CP_Mes,{M3},CP_Aprov,"<>Sim")', 'Compensações não aprovadas (pendentes ou recusadas)', 'Só contam com "Aprovado? = Sim".'),
    (f'=COUNTIFS(L_Alerta,"VINCULAR*",L_Mes,{M3})', '"Horas Compensadas" sem registro na aba Compensações', 'Registrar quem trabalhou, as horas e a aprovação.'),
    (f'=COUNTIFS(L_Alerta,"CLASSIFICAR*",L_Mes,{M3})+COUNTIFS(L_Alerta,"AJUSTAR*",L_Mes,{M3})', 'Atrasos a classificar / leve atraso como Horas Trabalhadas', 'Trocar o Status para "Entrada Atrasada… S/ ou C/ Justificativa".'),
    (f'=COUNTIFS(L_Alerta,"CADASTRO*",L_Mes,{M3})+COUNTIFS(L_Alerta,"DUPLICADO*",L_Mes,{M3})', 'Cadastro / vigência / lançamento duplicado', 'Corrigir Cadastro/Vigências ou apagar a linha repetida.'),
]
for i, (f, txt, oq) in enumerate(pend):
    rr = 13 + i
    cel(f'B{rr}', f, b=True)
    cel(f'C{rr}', txt, al='left'); RL.merge_cells(f'C{rr}:H{rr}')
    cel(f'I{rr}', oq, al='left'); RL.merge_cells(f'I{rr}:J{rr}')
    RL[f'N{rr}'] = f'=IF(B{rr}=0,1,"")'
    RL.conditional_formatting.add(f'B{rr}', FormulaRule(formula=[f'$B{rr}>0'], fill=fill(VERM), font=Font(color='FFC00000', bold=True)))

# --- ocorrências por tipo ---
faixa(21, 'SERRALHERIA — OCORRÊNCIAS POR TIPO')
cab(22, {'B': 'Qtde', 'C': 'Situação', 'E': 'Lançadas (h)', 'F': 'Ferista (h)', 'G': 'Compens. (h)', 'H': 'Desconta (h)',
         'I': 'O que significa', 'J': 'Valor (R$)'})
RL.merge_cells('C22:D22')
cats = [
    ('Injustificada', 'Faltas e atrasos sem justificativa', 'Falta ou atraso/saída sem documento. Desconta, menos o que o ferista, acordo ou compensação cobriu.'),
    ('Leve atraso', 'Leve atraso (10 a 30 min)', 'Atraso de 10 a 30 min pelo ponto. Desconta o tempo (status correto: Entrada Atrasada… S/ Justificativa).'),
    ('Férias', 'Férias', 'Desconta a jornada se o ferista não cobrir (férias têm prioridade na cobertura).'),
    ('Falta de Contratação', 'Posto vago (falta de contratação)', 'Vaga sem ninguém contratado. Desconta a jornada inteira.'),
    ('Falta de EPI', 'Falta de EPI', 'Não trabalhou por falta de EPI/uniforme. Desconta.'),
    ('Recesso', 'Recesso', 'Ausência por recesso. Conta no absenteísmo.'),
    ('Remanejamento', 'Remanejamento', 'Deslocado para outra atividade: o posto fica descoberto e desconta, salvo ferista.'),
    ('Situação atípica', 'Situação atípica', 'Situação fora do padrão (ex.: jogo da seleção). Horas não acordadas descontam.'),
    ('Outros', 'Outras (recusa, desligamento)', 'Recusa sem motivo legítimo, desligamento / pedido de demissão. Desconta.'),
]
cats_nao = [
    ('Justificada', 'Justificadas (atestado, comprovante, integração, afastamento)', 'Não desconta. Motivo na Observação (atestado: só "Dia Atestado").'),
    ('Treinamento', 'Treinamento', 'Não desconta.'),
]
def linha_cat(rr, cat, txt, sig):
    hcol = 'L_Ausencia' if cat in ('Justificada', 'Treinamento') else 'L_Bruto'
    crit = f'L_Mes,{M3},L_Equipe,P_EquipePrincipal,L_Categoria,"{cat}"'
    cel(f'B{rr}', f'=COUNTIFS({crit})', b=True)
    cel(f'C{rr}', txt, al='left'); RL.merge_cells(f'C{rr}:D{rr}')
    cel(f'E{rr}', f'=SUMIFS({hcol},{crit})', '[h]:mm')
    cel(f'F{rr}', f'=SUMIFS(L_Ferista,{crit})', '[h]:mm')
    cel(f'G{rr}', f'=SUMIFS(L_Compensado,{crit})', '[h]:mm')
    cel(f'H{rr}', f'=SUMIFS(L_HorasDesc,{crit})', '[h]:mm')
    cel(f'I{rr}', sig, al='left')
    cel(f'J{rr}', f'=SUMIFS(L_Valor,{crit})', '#,##0.00')
    RL[f'N{rr}'] = f'=IF(B{rr}=0,1,"")'
for i, (cat, txt, sig) in enumerate(cats):
    linha_cat(23 + i, cat, txt, sig)
r_vf, r_tot = 32, 33
cel(f'B{r_vf}', '')
cel(f'C{r_vf}', '( + ) Abatimento de Valor/Função (ferista de cargo menor que o coberto)', al='left'); RL.merge_cells(f'C{r_vf}:I{r_vf}')
cel(f'J{r_vf}', '=' + ix('AQ'), '#,##0.00')
RL[f'N{r_vf}'] = f'=IF(J{r_vf}=0,1,"")'
for col in 'BCDEFGHIJ':
    c = RL[f'{col}{r_tot}']; c.fill = fill(D9); c.font = t(9, True); c.border = BOX; c.alignment = Alignment(horizontal='center', vertical='center')
RL[f'B{r_tot}'] = '=SUM(B23:B31)'
RL[f'C{r_tot}'] = 'ABATIMENTO DO MÊS (resultado)'; RL[f'C{r_tot}'].alignment = Alignment(horizontal='left', vertical='center')
for col in 'EFGH':
    RL[f'{col}{r_tot}'] = f'=SUM({col}23:{col}31)'; RL[f'{col}{r_tot}'].number_format = '[h]:mm'
RL[f'J{r_tot}'] = f'=SUM(J23:J31)+J{r_vf}'; RL[f'J{r_tot}'].number_format = '#,##0.00'
RL[f'I{r_tot}'] = f'=IF(ROUND(J{r_tot}-D7,2)=0,"✔ confere com o Painel e o Resumo_Mensal","✘ diferença — conferir")'
RL.merge_cells(f'C{r_tot}:D{r_tot}')
cel(f'C34', 'NÃO DESCONTAM DO CONTRATO', b=True, al='left'); RL.merge_cells('C34:J34')
for i, (cat, txt, sig) in enumerate(cats_nao):
    linha_cat(35 + i, cat, txt, sig)
RL['N34'] = '=IF(AND(B35=0,B36=0),1,"")'

# --- ferista e compensação ---
faixa(38, 'FERISTA E COMPENSAÇÕES DO MÊS')
cab(39, {'B': 'Horas', 'C': 'Item', 'I': 'Leitura'})
RL.merge_cells('C39:H39'); RL.merge_cells('I39:J39')
itens = [
    ('=' + ix('V'), 'Ferista: horas trabalhadas', 'Horas do ferista no mês (ele nunca gera absenteísmo próprio).'),
    ('=' + ix('D'), 'Ferista: horas que cobriram faltas', 'Cascata da EI32: férias primeiro, depois a maior falta; a sobra (≥ 1h) é doada à próxima.'),
    ('=' + ix('W'), 'Ferista: horas sem cobertura', 'Horas do ferista que não tinham falta para cobrir.'),
    ('=' + ix('AJ'), 'Compensações aprovadas no mês', 'Só "Aprovado? = Sim" na aba Compensações.'),
    ('=' + ix('AK'), 'Adicional não solicitado', 'Pessoa a mais, não pedida: abate 1:1 depois da compensação.'),
    ('=' + ix('E'), 'Compensação + adicional usados (abatem em bloco)', 'Abatem o absenteísmo do mês 1:1 (valor proporcional).'),
    ('=' + ix('AO'), 'Sobra de compensação (horas positivas)', 'Excedente para acerto comercial; não vira crédito.'),
]
for i, (f, txt, lei) in enumerate(itens):
    rr = 40 + i
    cel(f'B{rr}', f, '[h]:mm', b=True)
    cel(f'C{rr}', txt, al='left'); RL.merge_cells(f'C{rr}:H{rr}')
    cel(f'I{rr}', lei, al='left'); RL.merge_cells(f'I{rr}:J{rr}')
RL['N40'] = '=IF(AND(B40=0,B41=0,B42=0),1,"")'; RL['N41'] = '=N40'; RL['N42'] = '=N40'
cel('C47', '="Abatimento de Valor/Função no mês: R$ "&FIXED(' + ix('AQ') + ',2)&" (detalhe na coluna Explicação abaixo)."', al='left')
RL.merge_cells('C47:J47'); RL['N47'] = f'=IF({ix("AQ")}=0,1,"")'

# --- detalhe dia a dia ---
D0, DN = 50, 300
faixa(49, 'SERRALHERIA — DETALHE DIA A DIA')
cab(D0, {'B': 'Data', 'C': 'Funcionário', 'D': 'Situação', 'E': 'Lançadas (h)', 'F': 'Ferista (h)', 'G': 'Compens. (h)',
         'H': 'Desconta (h)', 'I': 'Explicação', 'J': 'Valor (R$)'})
for k in range(1, DN + 1):
    rr = D0 + k
    RL[f'L{rr}'] = f'=IFERROR(MATCH({k},L_RelN,0),0)'
    m = f'$L{rr}'
    cel(f'B{rr}', f'=IF({m}=0,"",INDEX(L_Data,{m}))', 'dd/mm/yyyy')
    cel(f'C{rr}', f'=IF({m}=0,"",IF(INDEX(L_Nome,{m})="","Posto vago — "&INDEX(L_Funcao,{m}),INDEX(L_Nome,{m})))', al='left')
    cel(f'D{rr}', f'=IF({m}=0,"",IF(INDEX(L_Categoria,{m})="Leve atraso","Leve atraso (10 a 30 min)",IF(INDEX(Lancamentos!$J$4:$J${LAST},{m})<>"","Situação atípica",'
                  f'IFERROR(INDEX(S_Rotulo,MATCH(INDEX(L_Status,{m}),S_Status,0)),INDEX(L_Status,{m})))))', al='left')
    cel(f'E{rr}', f'=IF({m}=0,"",INDEX(L_Bruto,{m})+IF(INDEX(L_Categoria,{m})="Justificada",INDEX(L_Ausencia,{m}),0))', '[h]:mm')
    cel(f'F{rr}', f'=IF({m}=0,"",INDEX(L_Ferista,{m}))', '[h]:mm')
    cel(f'G{rr}', f'=IF({m}=0,"",INDEX(L_Compensado,{m}))', '[h]:mm')
    cel(f'H{rr}', f'=IF({m}=0,"",INDEX(L_HorasDesc,{m}))', '[h]:mm')
    cel(f'I{rr}', f'=IF({m}=0,"",TRIM(INDEX(L_Explica,{m})&IF(INDEX(L_Obs,{m})<>""," Obs.: "&INDEX(L_Obs,{m}),"")'
                  f'&IF(AND(INDEX(L_Explica,{m})="",INDEX(L_HorasDesc,{m})=0)," Não desconta.","")'
                  f'&IF(AND(INDEX(L_Div,{m})<>"",INDEX(L_Conf,{m})<>"Sim")," ⚠ PONTO: "&INDEX(L_Div,{m}),"")))', al='left', sz=8)
    cel(f'J{rr}', f'=IF({m}=0,"",INDEX(L_Valor,{m})+INDEX(L_ValorFuncao,{m}))', '#,##0.00')
    RL[f'N{rr}'] = f'=IF({m}=0,1,"")'
    RL.row_dimensions[rr].height = 24
rdet = D0 + DN + 1
RL[f'B{rdet}'] = f'=IF(COUNTIF(L_Rel,1)>{DN},"⚠ Mais de {DN} ocorrências no mês: o detalhe mostra as {DN} primeiras.","")'
RL[f'B{rdet}'].font = t(8, True, 'FFC00000'); RL[f'N{rdet}'] = f'=IF(COUNTIF(L_Rel,1)>{DN},"",1)'

# --- horas extras ---
H0, HN = rdet + 2, 60
faixa(H0, 'HORAS EXTRAS DO MÊS — CASO A CASO')
cab(H0 + 1, {'B': 'Data', 'C': 'Funcionário', 'D': 'Tratamento', 'E': 'Tipo do dia', 'F': 'Líquidas (h)', 'G': 'HE 50% (h)',
             'H': 'HE 100% (h)', 'I': 'Alerta / descrição', 'J': 'Valor (R$)'})
RL[f'N{H0}'] = f'=IF(COUNTIF(H_RelN,">0")=0,1,"")'; RL[f'N{H0 + 1}'] = f'=N{H0}'
for k in range(1, HN + 1):
    rr = H0 + 1 + k
    RL[f'L{rr}'] = f'=IFERROR(MATCH({k},H_RelN,0),0)'
    m = f'$L{rr}'
    cel(f'B{rr}', f'=IF({m}=0,"",INDEX(H_Data,{m}))', 'dd/mm/yyyy')
    cel(f'C{rr}', f'=IF({m}=0,"",INDEX(H_Nome,{m}))', al='left')
    cel(f'D{rr}', f'=IF({m}=0,"",IF(INDEX(H_Trat,{m})="","SEM TRATAMENTO",INDEX(H_Trat,{m})))', al='left')
    cel(f'E{rr}', f'=IF({m}=0,"",INDEX(H_Tipo,{m}))')
    cel(f'F{rr}', f'=IF({m}=0,"",INDEX(H_Liq,{m}))', '[h]:mm')
    cel(f'G{rr}', f'=IF({m}=0,"",INDEX(H_50,{m}))', '[h]:mm')
    cel(f'H{rr}', f'=IF({m}=0,"",INDEX(H_100,{m}))', '[h]:mm')
    cel(f'I{rr}', f'=IF({m}=0,"",TRIM(INDEX(H_Alerta,{m})&" "&INDEX(Lanc_HE!$J$4:$J${HLAST},{m})))', al='left', sz=8)
    cel(f'J{rr}', f'=IF({m}=0,"",INDEX(H_Valor,{m}))', '#,##0.00')
    RL[f'N{rr}'] = f'=IF({m}=0,1,"")'
rtot_he = H0 + 2 + HN
cel(f'C{rtot_he}', 'TOTAL HE DO MÊS', b=True, al='left')
cel(f'F{rtot_he}', f'=SUMIFS(H_Liq,H_Mes,{M3})', '[h]:mm', b=True)
cel(f'J{rtot_he}', f'=SUMIFS(H_Valor,H_Mes,{M3})', '#,##0.00', b=True)
RL[f'N{rtot_he}'] = f'=N{H0}'

# --- compensações ---
C0, CN = rtot_he + 2, 30
faixa(C0, 'COMPENSAÇÕES DO MÊS')
cab(C0 + 1, {'B': 'Data', 'C': 'Quem trabalhou', 'D': 'Compensa a falta de', 'E': 'Horas', 'F': 'Aprovado?', 'G': 'Contam (h)',
             'I': 'Situação / observação'})
RL.merge_cells(f'G{C0 + 1}:H{C0 + 1}'); RL.merge_cells(f'I{C0 + 1}:J{C0 + 1}')
RL[f'N{C0}'] = f'=IF(COUNTIF(CP_RelN,">0")=0,1,"")'; RL[f'N{C0 + 1}'] = f'=N{C0}'
for k in range(1, CN + 1):
    rr = C0 + 1 + k
    RL[f'L{rr}'] = f'=IFERROR(MATCH({k},CP_RelN,0),0)'
    m = f'$L{rr}'
    cel(f'B{rr}', f'=IF({m}=0,"",INDEX(CP_Data,{m}))', 'dd/mm/yyyy')
    cel(f'C{rr}', f'=IF({m}=0,"",INDEX(CP_Quem,{m}))', al='left')
    cel(f'D{rr}', f'=IF({m}=0,"",IF(INDEX(CP_Falta,{m})="","—",INDEX(CP_Falta,{m})))', al='left')
    cel(f'E{rr}', f'=IF({m}=0,"",INDEX(CP_Horas,{m}))', '[h]:mm')
    cel(f'F{rr}', f'=IF({m}=0,"",INDEX(CP_Aprov,{m}))')
    cel(f'G{rr}', f'=IF({m}=0,"",INDEX(CP_Conta,{m}))', '[h]:mm'); RL.merge_cells(f'G{rr}:H{rr}')
    cel(f'I{rr}', f'=IF({m}=0,"",TRIM(INDEX(CP_Situacao,{m})&" "&INDEX(CP_Obs,{m})))', al='left', sz=8); RL.merge_cells(f'I{rr}:J{rr}')
    RL[f'N{rr}'] = f'=IF({m}=0,1,"")'
FIM = C0 + 1 + CN
RL['L1'] = FIM
for col in 'LMN':
    RL.column_dimensions[col].hidden = True
RL.print_area = f'B1:J{FIM}'
RL.page_setup.orientation = 'landscape'
RL.page_setup.paperSize = RL.PAPERSIZE_A4
RL.page_setup.fitToWidth = 1; RL.page_setup.fitToHeight = 0
RL.sheet_properties.pageSetUpPr.fitToPage = True
RL.print_options.horizontalCentered = True
RL.page_margins.left = RL.page_margins.right = 0.4
RL.page_margins.top = RL.page_margins.bottom = 0.5
RL.print_title_rows = f'{D0}:{D0}'
RL.freeze_panes = 'A5'
RL.sheet_properties.tabColor = TK
from openpyxl.worksheet.pagebreak import Break
RL.row_breaks.append(Break(id=48))

# =====================================================================================
# 10. LEIA-ME (textos que mudaram)
# =====================================================================================
LM = wb['LEIA-ME']
txt = {
    'B6': ('Cada desconto é montado na própria linha de Lançamentos, da esquerda para a direita:  Horas brutas  −  coberto pelo ferista  −  '
           'perdoado por acordo  −  compensado (em bloco, proporcional no mês)  =  Horas descontadas  ×  Valor HH  =  R$. '
           'A coluna AG ("De onde vem") explica isso em texto em cada linha. Regras iguais às da EI32 (padrão thyssenkrupp).'),
    'B9': 'Cabeçalho azul / célula cinza-claro', 'C9': 'Você preenche.',
    'B10': 'Cabeçalho marinho',
    'B13': 'Vermelho-claro', 'C13': 'Alerta: precisa da sua atenção (ponto divergente, valor/função, pendências).',
    'C16': 'Escolha o mês em C3. Mostra contrato, abatimento, valor final, de onde vem cada real descontado e a lista de conferência. Linha 2 = atalhos; botões "Atualizar Pontos" e "Criar Relatório (PDF)" ficam na linha 3.',
    'C18': 'Quem trabalhou para compensar e quantas horas. Só abate quando "Aprovado? = Sim". Abate em bloco no mês (igual EI32).',
    'C30': 'Coluna AH (alerta), o Painel de Verificação e a aba Conferencia_Ponto (depois de "Atualizar Pontos").',
    'B33': 'Exemplo: Giovani faltou em 10/09 (desconto de 9h). Ele trabalhou no sábado 12/09 para compensar.',
    'C35': 'Data em que trabalhou = 12/09 · Quem trabalhou = Giovani · Horas = 8:00 · Aprovado? = Pendente. (De quem é a falta e a data são opcionais.)',
    'C36': 'Mude para "Sim". As 8h aprovadas abatem o absenteísmo do mês da equipe em bloco (1:1), proporcional em todas as linhas com desconto.',
    'C37': 'Igual: registre quem trabalhou e as horas. Vale para autocompensação, colega cobrindo colega e adicional não solicitado (status próprio).',
    'C38': ('Automático (regra EI32): no dia em que o ferista trabalha, as horas dele cobrem as faltas dos serralheiros '
            '(nunca do encarregado): FÉRIAS PRIMEIRO, depois a MAIOR falta. A sobra é doada à próxima falta se for de pelo menos 1h. '
            'Se o ferista tiver cargo menor que o coberto, desconta a diferença de valor-hora (Abatimento de Valor/Função, col. BA).'),
    'C39': 'Não estiver aprovado. A coluna "Situação" da aba Compensações diz qual é o caso. Sobra de compensação vira "horas positivas" (acerto comercial).',
}
for a, v in txt.items():
    LM[a] = v
LM['B26'] = 'ROTINA DIÁRIA'
extra = [('B61', 'PONTOS E RELATÓRIO (suplemento Controle de Horas TK)', True),
         ('B62', 'Atualizar Pontos', False), ('C62', 'Lê a base de pontos (Pontos 2026.xlsx, aba "Base unificada"), preenche o ponto (col. D), marca divergências em BC e cria as linhas de dias com ponto sem lançamento. Confira na aba Conferencia_Ponto e marque "Conferido? = Sim".', False),
         ('B63', 'Criar Relatório (PDF)', False), ('C63', 'Gera o PDF do mês escolhido no Painel em Relatorios\\03 - EI36\\<ano>\\<mm - Mês> - EI36.pdf (um por mês; o mesmo mês é substituído).', False),
         ('B64', 'Instalar botões', False), ('C64', 'Uma vez por planilha: Suplementos › Controle de Horas TK › Instalar botões nesta planilha › Ctrl + S.', False)]
for a, v, tit in extra:
    LM[a] = v
    src = LM['B54'] if tit else (LM['B55'] if a.startswith('B') else LM['C55'])
    LM[a].font = src.font.copy(); LM[a].fill = src.fill.copy(); LM[a].alignment = src.alignment.copy(); LM[a].border = src.border.copy()

# =====================================================================================
# 11. VISUAL — paleta da EI32 (sem mudar posição de nada)
# =====================================================================================
FILL_MAP = {'FF2F5597': TK, 'FFD9E1F2': F2, 'FFDDEBF7': D9, 'FFFFF2CC': F2}
FONT_MAP = {'FF000000': TXT, 'FF0000FF': TXT, 'FF595959': TXT}
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    maxr = ws.max_row
    for row in ws.iter_rows(min_row=1, max_row=maxr):
        for c in row:
            if c.fill is not None and c.fill.fill_type == 'solid':
                rgb = c.fill.fgColor.rgb
                if rgb in FILL_MAP:
                    c.fill = fill(FILL_MAP[rgb])
                elif rgb == 'FF1F3864' and c.row == 1:
                    c.fill = PatternFill(fill_type=None)
            fo = c.font
            if fo is not None:
                col = fo.color.rgb if (fo.color is not None and isinstance(fo.color.rgb, str)) else None
                if c.row == 1 and ws.title not in ('Relatorio', 'Conferencia_Ponto') and c.value is not None and col == 'FFFFFFFF':
                    c.font = Font(name='Arial', size=16, bold=True, color=TK)
                elif col in FONT_MAP or (col is None and c.value is not None and fo.name == 'Arial'):
                    c.font = Font(name='Arial', size=fo.sz, bold=fo.b, italic=fo.i, underline=fo.u, color=FONT_MAP.get(col, TXT))
# texto branco em toda faixa azul TK / marinho
for ws in wb.worksheets:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 120)):
        for c in row:
            if c.fill is not None and c.fill.fill_type == 'solid' and c.fill.fgColor.rgb in (TK, NAVY) and c.value is not None:
                fo = c.font
                c.font = Font(name='Arial', size=fo.sz or 9, bold=True, italic=fo.i, color='FFFFFFFF')
# cabeçalhos das abas de lançamento / cadastro: azul TK = você preenche; marinho = automático; aux cinza
def cabecalho(ws, hr, ultima, r_ini, r_fim):
    for ci in range(1, CI(ultima) + 1):
        h = ws.cell(hr, ci)
        if h.value is None:
            continue
        aux = 'aux' in str(h.value).lower() or 'vig#' in str(h.value).lower()
        vals = [ws.cell(r, ci).value for r in range(r_ini, min(r_fim, r_ini + 60))]
        formula = any(isinstance(v, str) and v.startswith('=') for v in vals)
        if aux:
            h.fill = fill('FFEDEDED'); h.font = Font(name='Arial', size=9, bold=True, color=CINZA)
        else:
            h.fill = fill(NAVY if formula else TK); h.font = HDR_FONT
        h.alignment = HDR_AL; h.border = BOXH
        if not formula and not aux:
            for r in range(r_ini, r_fim + 1):
                x = ws.cell(r, ci)
                x.fill = fill(F2)
    ws.row_dimensions[hr].height = max(ws.row_dimensions[hr].height or 0, 40)

cabecalho(LW, 3, 'BH', 4, LAST)
cabecalho(HW, 3, 'AH', 4, HLAST)
cabecalho(CW, 3, 'Q', 4, CLAST)
cabecalho(CAD, 3, 'L', 4, 200)
cabecalho(wb['Vigencias'], 3, 'I', 4, 300)
for ws, a in ((LW, 'A2'), (HW, 'A2'), (CW, 'A2'), (CAD, 'A2'), (wb['Vigencias'], 'A2'), (RW, 'A2'), (T, 'A2')):
    ws[a].font = Font(name='Arial', size=9, italic=True, color=CINZA)
# Painel: números grandes em azul TK (como a Gestão da EI32)
for a in ('B6', 'D6', 'F6', 'H6'):
    c = P[a]; c.font = Font(name='Arial', size=18, bold=True, color=TK); c.fill = fill('FFFFFFFF')
for a in ('B5', 'D5', 'F5', 'H5'):
    P[a].font = Font(name='Arial', size=9, bold=True, color=TXT)
P['C3'].fill = fill(D9); P['C3'].font = Font(name='Arial', size=12, bold=True, color=TXT)
# Tabelas: cabeçalhos de bloco (linha 3, A20, A35, E14) em azul TK
for row in T.iter_rows(min_row=1, max_row=60):
    for c in row:
        if c.fill is not None and c.fill.fill_type == 'solid' and c.fill.fgColor.rgb in ('FFD9D9D9',) and c.font.b:
            c.fill = fill(TK); c.font = HDR_FONT
for a in ('A35', 'AF3', 'AG3', 'AH3'):
    T[a].fill = fill(TK); T[a].font = HDR_FONT; T[a].alignment = HDR_AL
for a in ('A36', 'A37', 'A38', 'A18', 'A19'):
    T[a].font = Font(name='Arial', size=9, color=TXT); T[a].border = BOX
for a in ('B36', 'B37', 'B38', 'B18', 'B19'):
    T[a].fill = fill(F2); T[a].font = Font(name='Arial', size=9, color=TXT); T[a].border = BOX
for a in ('C36', 'C37', 'C38', 'C18', 'C19'):
    T[a].font = Font(name='Arial', size=8, italic=True, color=CINZA); T[a].border = BOX
# formatação condicional: cores da EI32 (verde / vermelho claro)
DXF_MAP = {'FFE2EFDA': VERDE, 'FFFCE4D6': VERM, 'FFFFE699': VERM}
for d in wb._differential_styles.dxf:
    if d.fill is not None and d.fill.bgColor is not None and d.fill.bgColor.rgb in DXF_MAP:
        novo = DXF_MAP[d.fill.bgColor.rgb]
        d.fill = PatternFill(bgColor=novo, fgColor=novo)
# guias por grupo (igual EI32)
abas = {'LEIA-ME': 'FF7FC4E8', 'Painel': TK, 'Relatorio': TK, 'Resumo_Mensal': TK, 'Por_Colaborador': TK,
        'Lancamentos': NAVY, 'Compensacoes': NAVY, 'Conferencia_Ponto': NAVY, 'Lanc_HE': NAVY,
        'Cadastro': 'FF8C8C8C', 'Vigencias': 'FF8C8C8C', 'Tabelas': 'FF8C8C8C'}
for n, cor in abas.items():
    wb[n].sheet_properties.tabColor = cor
LW.freeze_panes = 'C4'
wb.active = wb.sheetnames.index('Painel')
for ws in wb.worksheets:
    ws.sheet_view.tabSelected = (ws.title == 'Painel')

for row in P.iter_rows(min_row=3, max_row=40, min_col=CI('K'), max_col=CI('Q')):
    for c in row:
        fo = c.font
        cor = fo.color.rgb if (fo.color is not None and isinstance(fo.color.rgb, str)) else TXT
        c.font = Font(name='Arial', size=9, bold=fo.b, italic=fo.i, color=cor)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('ok', OUT)
