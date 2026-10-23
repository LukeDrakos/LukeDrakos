"""Padronização visual e de orientação das planilhas EI32 e EI36 (mesmo guia, mesmas dicas, mesmo visual).
Não altera nenhuma fórmula de cálculo.

uso: python padronizar.py EI32|EI36 entrada.xlsx saida.xlsx [logo.png]
"""
import sys, copy
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import column_index_from_string as CI, get_column_letter as CL
from openpyxl.worksheet.cell_range import CellRange

LAY, SRC, OUT = sys.argv[1:4]
LOGO = sys.argv[4] if len(sys.argv) > 4 else None
wb = openpyxl.load_workbook(SRC)

TK, NAVY, TXT, CINZA, BORDA = 'FF0095D8', 'FF1F3864', 'FF4D4D4D', 'FF7F7F7F', 'FFD9D9D9'
F2, D9, AUXF = 'FFF2F2F2', 'FFD9D9D9', 'FFEDEDED'
fill = lambda c: PatternFill('solid', fgColor=c)
fino = Side(style='thin', color=BORDA)
BOX = Border(left=fino, right=fino, top=fino, bottom=fino)
branco = Side(style='thin', color='FFFFFFFF')
BOXH = Border(left=branco, right=branco, top=branco, bottom=branco)

# ======================================================================================
# Layout de cada planilha (só o que muda entre EI32 e EI36)
# ======================================================================================
if LAY == 'EI32':
    L = dict(cod='EI32', equipes='Elétrica e Climatização', painel='Gestao_Contrato', mes='B3', lanc='Lanc_HORAS',
             l_hdr=1, l_ini=2, l_fim=3873, h_hdr=1, h_ini=2, h_fim=693, c_ini=4, c_fim=503,
             cad_hdr=2, cad_ini=3, cad_fim=150, vig_hdr=2, vig_ini=3, vig_fim=149,
             data='A', nome='E', status='J', ponto='Y', obs='X', posto='R', turno='S', atip='L', atip_desc='M',
             atip_just='N', acordo='P', aus_man='M', trab_man=None, explica='AA', alerta='Z', div='AV', ok='AW',
             valfun='BA', resumo='Analises', sla='aba SLA', relat='Relatorio',
             he=dict(data='A', func='C', nome='E', horario='M', fer='N', horas='P', trat='V', solic='Y', area='Z',
                     desc='AA', ponto='AB', receb='AC', req='AD', div='AH', ok='AI', alerta='AE'),
             tabelas_status='Tabelas A:C', contrato='Tabelas U:W', param_pontos='Tabelas K16',
             pontos_desde=None, comp_regra=('As horas aprovadas do mês quitam primeiro as horas pendentes (coluna Q) de cada equipe; '
                                            'a sobra da Elétrica cobre a Climatização; o resto abate o absenteísmo do mês 1:1.'))
else:
    L = dict(cod='EI36', equipes='Serralheria', painel='Painel', mes='C3', lanc='Lancamentos',
             l_hdr=3, l_ini=4, l_fim=3489, h_hdr=3, h_ini=4, h_fim=1000, c_ini=4, c_fim=999,
             cad_hdr=3, cad_ini=4, cad_fim=200, vig_hdr=3, vig_ini=4, vig_fim=300,
             data='A', nome='B', status='C', ponto='D', obs='G', posto='H', turno='I', atip='J', atip_desc='K',
             atip_just='L', acordo='M', aus_man='E', trab_man='F', explica='AG', alerta='AH', div='BC', ok='BD',
             valfun='BA', resumo='Resumo_Mensal', sla='Resumo_Mensal (coluna I)', relat='Relatorio',
             he=dict(data='A', nome='B', func='C', horario='D', fer='E', horas='F', trat='G', solic='H', area='I',
                     desc='J', ponto='K', receb='L', req='M', div='AE', ok='AF', alerta='AB'),
             tabelas_status='Tabelas O:U', contrato='Tabelas A21:B32', param_pontos='Tabelas B18',
             pontos_desde='Tabelas B38', comp_regra=('As horas aprovadas do mês abatem o absenteísmo da equipe 1:1, em bloco '
                                                     '(proporcional em todas as linhas com desconto). A sobra vira "horas positivas".'))
Lh = L['he']

# ======================================================================================
# 1. Dicas: mensagem ao clicar (validação) + balão no cabeçalho (comentário)
# ======================================================================================
def prompts_lanc():
    p = {
        L['data']: ('Data', 'Data do dia (dd/mm/aaaa). Uma linha por pessoa por dia. Fim de semana/feriado trabalhado vai na aba Lanc_HE.'),
        L['nome']: ('Funcionário', 'Escolha pela lista (nome igual ao Cadastro). Deixe VAZIO para posto vago e informe a função do posto na coluna ' + L['posto'] + '.'),
        L['status']: ('Status do dia', 'O que aconteceu no dia. É o Status que decide se desconta. Atraso/saída sem documento (inclusive leve, 10–30 min) = "Entrada Atrasada ou Saída Antecipada S/ Justificativa".'),
        L['ponto']: ('Ponto do dia', 'Cole como vem do relógio: entrada#saída = total (ex.: 07:58#17:21 = 08:53). Total = última − primeira − 30 min. O botão Atualizar Pontos preenche sozinho.'),
        L['obs']: ('Observação', 'Motivo de toda justificada, remanejamento, treinamento, atípica, acordo ou ajuste. Atestado: escreva só "Dia Atestado" (sem motivo médico).'),
        L['posto']: ('Função do posto vago', 'Só em POSTO VAGO (funcionário vazio): a função do posto que ficou sem ninguém. Desconta a jornada inteira.'),
        L['turno']: ('Turno no dia', 'Só se a pessoa trabalhou num turno diferente do cadastro NESTE dia. Muda a jornada só desta linha.'),
        L['atip']: ('Situação atípica', 'Só para evento fora do padrão (ex.: jogo da seleção). Presente = trabalhou parte do dia; Saída = saiu. Informe as horas nas colunas ao lado.'),
        L['acordo']: ('Horas acordadas', 'Horas perdoadas por acordo comercial com o cliente (formato h:mm). Saem do desconto. Escreva o acordo na Observação.'),
        L['div']: ('Divergência do ponto', 'Preenchida pelo botão Atualizar Pontos. Não apague: analise e marque "Conferido? = Sim" na coluna ao lado.'),
        L['ok']: ('Conferido?', 'Depois de analisar a divergência do ponto, marque Sim. A linha deixa de ficar vermelha.'),
    }
    if LAY == 'EI32':
        p.update({
            'K': ('Não usar', 'Coluna antiga (Horas Transferidas). Não preencha.'),
            'M': ('Absenteísmo manual', 'Horas de ausência que o cálculo automático não pegou (h:mm). Desconta. Explique na Observação.'),
            'N': ('Justificadas manual', 'Horas justificadas lançadas à mão (h:mm). Não descontam. Explique na Observação.'),
            'O': ('Cobertura manual ferista', 'Só se o ferista cobriu ESTE posto de forma diferente da automática (h:mm). Nunca para encarregado.'),
            'Q': ('Horas pendentes', 'Dívida de horas combinada (ex.: Copa). Não desconta agora; é quitada pelas compensações aprovadas.'),
        })
    else:
        p.update({
            'E': ('Horas ausente (manual)', 'Só se NÃO houver ponto: horas que a pessoa ficou fora (h:mm). Com ponto, deixe vazio: o cálculo é automático.'),
            'F': ('Horas trabalhadas (manual)', 'Só para trabalho fora da jornada (ex.: sábado de compensação de quem é ADM), em h:mm.'),
            'K': ('Desconto extra (h)', 'Situação atípica: horas que DESCONTAM (h:mm). Escreva o motivo na Observação.'),
            'L': ('Justificadas extra (h)', 'Situação atípica: horas justificadas que NÃO descontam (h:mm). Motivo na Observação.'),
        })
    return p


def comentarios_lanc():
    c = {
        L['status']: 'O Status decide se a linha desconta. Veja a tabela "COMO LANÇAR CADA SITUAÇÃO" na aba Como_Usar.',
        L['ponto']: 'Ponto do dia: entrada#saída = total. Total = última marcação − primeira − 30 min (almoço). Preenchido pelo botão Atualizar Pontos.',
        L['obs']: 'Motivo obrigatório em justificadas, remanejamento, treinamento, situação atípica, acordo, ajuste manual, recusa, afastamento e desligamento. Atestado: só "Dia Atestado".',
        L['posto']: 'Posto vago: Funcionário vazio + a função do posto aqui. Desconta a jornada inteira.',
        L['turno']: 'Turno diferente só neste dia. Se ADM entrar antes das 07:00 e sair até 14:00, o Turno 2 é automático.',
        L['acordo']: 'Horas perdoadas por acordo comercial. Não descontam.',
        L['explica']: 'AUTOMÁTICO. Explica em texto a linha: ferista, compensação, acordo, Valor/Função e quanto desconta.',
        L['alerta']: 'AUTOMÁTICO. O que precisa ser feito nesta linha (CONFERIR, CLASSIFICAR, AJUSTAR STATUS, VINCULAR…). Vazio = linha certa.',
        L['div']: 'Preenchida pelo botão Atualizar Pontos. Enquanto houver texto aqui e "Conferido?" não for Sim, a linha fica vermelha.',
        L['ok']: 'Marque Sim depois de analisar a divergência do ponto.',
        L['valfun']: 'AUTOMÁTICO. Ferista de cargo menor cobrindo cargo maior: (HH coberto − HH ferista) × horas cobertas. Ferista igual ou maior = zero.',
    }
    if LAY == 'EI32':
        c.update({
            'O': 'Só para cobertura MANUAL do ferista. A normal é automática: férias primeiro, depois a maior falta; sobra ≥ 1h é doada à próxima falta. Nunca cobre encarregado.',
            'Q': 'Horas que a equipe ficou devendo (ex.: Copa). Não descontam agora; são quitadas pelas compensações aprovadas.',
            'V': 'AUTOMÁTICO. Horas de ausência: jornada − ponto (atrasos/saídas) ou jornada inteira (faltas).',
            'W': 'AUTOMÁTICO. Valor descontado = horas que descontam × Valor HH (já sem ferista, acordo e pendências).',
            'AC': 'AUTOMÁTICO. Horas cobertas pelo ferista neste posto (cascata: férias primeiro, depois a maior falta).',
            'M': 'Horas de absenteísmo lançadas à mão (descontam).',
            'N': 'Horas justificadas lançadas à mão (não descontam).',
        })
    else:
        c.update({
            'E': 'Só sem ponto: horas que a pessoa ficou ausente. Com ponto, deixe vazio.',
            'F': 'Só para trabalho fora da jornada (ex.: sábado de compensação de quem é ADM).',
            'J': 'Situação atípica: escolha Presente ou Saída e informe as horas em K (desconta) ou L (justifica).',
            'Y': 'AUTOMÁTICO. Horas brutas a descontar (ausência que conta no absenteísmo + desconto extra).',
            'Z': 'AUTOMÁTICO. Horas cobertas pelo ferista (cascata da EI32: férias primeiro, depois a maior falta; sobra ≥ 1h é doada).',
            'AA': 'AUTOMÁTICO. Compensação do mês em bloco (proporcional). Vem das horas aprovadas na aba Compensações.',
            'AB': 'AUTOMÁTICO. Horas perdoadas por acordo (coluna M).',
            'AC': 'AUTOMÁTICO. Horas que descontam = brutas − ferista − acordo − compensação.',
            'AD': 'AUTOMÁTICO. Valor descontado = horas que descontam × Valor HH.',
        })
    return c


def prompts_he():
    return {
        Lh['data']: ('Data', 'Data em que a hora extra foi feita (dd/mm/aaaa).'),
        Lh['nome']: ('Funcionário', 'Escolha pela lista.'),
        Lh['func']: ('Função solicitada', 'Só se a HE foi pedida numa função DIFERENTE da pessoa. O valor segue a função solicitada. Senão, deixe vazio.'),
        Lh['horario']: ('Horário solicitado', 'Horário pedido pela área (ex.: 08:00 às 12:00).'),
        Lh['fer']: ('É feriado?', 'Sim se a data for feriado (HE 100%). Sábado: primeiras 4h a 100%, resto 50%. Domingo/feriado: 100%.'),
        Lh['horas']: ('Horas totais', 'Total de horas no dia (h:mm), sem tirar almoço: o desconto de 1h no fim de semana/feriado é automático.'),
        Lh['trat']: ('Tratamento (obrigatório)', 'Pagas por EI / pela área = valor. HORAS COMPENSADAS = registre também na aba Compensações. Remanejado = só registro. Vazio = pendência.'),
        Lh['solic']: ('Solicitante', 'Quem pediu a hora extra.'),
        Lh['area']: ('Área', 'Área que pediu a hora extra.'),
        Lh['desc']: ('Descrição', 'O que foi feito na hora extra.'),
        Lh['ponto']: ('Ponto do dia', 'Cole o ponto: entrada#saída = total. O botão Atualizar Pontos confere e corrige as horas.'),
        Lh['receb']: ('Recebeu o ponto?', 'Sim quando o ponto do dia já foi recebido e conferido.'),
        Lh['req']: ('Requisição / pedido', 'Número da requisição de compra ou pedido da HE paga por EI.'),
        Lh['div']: ('Divergência do ponto', 'Preenchida pelo botão Atualizar Pontos. Analise e marque "Conferido? = Sim" na coluna ao lado.'),
        Lh['ok']: ('Conferido?', 'Depois de analisar a divergência, marque Sim. A linha deixa de ficar vermelha.'),
    }


def comentarios_he():
    return {
        Lh['trat']: 'OBRIGATÓRIO. PAGAS POR EI (fora do contrato, entra no saldo) · PAGAS PELA ÁREA · HORAS COMPENSADAS (registrar na aba Compensações) · REMANEJADO P/ OUTRA ÁREA. Vazio = pendência.',
        Lh['horas']: 'Horas totais do dia. Sábado/domingo/feriado acima do limite descontam 1h de almoço sozinhos.',
        Lh['func']: 'Preencha só se a área pediu outra função: o valor da HE segue a função solicitada.',
        Lh['ponto']: 'Ponto do dia (entrada#saída = total). Conferido pelo botão Atualizar Pontos.',
        Lh['div']: 'Preenchida pelo botão Atualizar Pontos. Vermelho até "Conferido? = Sim".',
        Lh['alerta']: 'AUTOMÁTICO. O que falta nesta HE (tratamento, registro da compensação, ponto). Vazio = certo.',
    }


PROMPT_COMP = {
    'A': ('Data em que trabalhou', 'Dia em que a pessoa trabalhou para compensar (ex.: o sábado).'),
    'B': ('Quem trabalhou', 'Escolha pela lista.'),
    'C': ('Horas', 'Horas líquidas do dia (h:mm). Use as horas líquidas da HE ou do ponto.'),
    'D': ('Falta de (opcional)', 'Só para rastreio: de quem era a falta que esta compensação cobre.'),
    'E': ('Data da falta (opcional)', 'Só para rastreio: dia da falta compensada.'),
    'F': ('Aprovado?', 'Sim = conta na compensação do mês. Pendente = ainda não conta. Não = recusada.'),
    'G': ('Aprovado por / obs.', 'Quem aprovou e qualquer detalhe que ajude a explicar depois.'),
}
COM_COMP = {
    'A': 'Uma linha por dia trabalhado para compensar. Só conta com "Aprovado? = Sim". ' + L['comp_regra'],
    'F': 'Só "Sim" abate. Pendente aparece no Painel de Verificação até ser aprovada.',
    'L': 'AUTOMÁTICO. Diz o que falta (PREENCHER / CONFERIR / PENDENTE) ou OK.',
}

if LAY == 'EI32':
    PROMPT_CAD = {'A': ('Registro', 'Matrícula da pessoa (igual ao relógio de ponto). É por ela que a macro de pontos acha a pessoa.'),
                  'B': ('Funcionário', 'Nome completo, igual ao usado nos lançamentos.'),
                  'C': ('Estilo do contrato', 'Normal / Fixo, Ferista, SPOT, Compensação ou Desligado.'),
                  'D': ('Equipe', 'Elétrica ou Climatização.'),
                  'E': ('Função', 'Função atual. O Valor HH vem da tabela de funções (aba Tabelas).'),
                  'F': ('Turno', 'ADM, TURNO 1, 2 ou 3.'),
                  'J': ('Status', 'Ativo ou Desligado. Não apague quem saiu.'),
                  'K': ('Observação', 'Ex.: data de adesão ou de desligamento.')}
    PROMPT_VIG = {'A': ('Registro', 'Mesmo registro do Cadastro.'),
                  'B': ('Funcionário', 'Escolha pela lista.'),
                  'C': ('Equipe', 'Equipe a partir desta data.'),
                  'D': ('Início', 'Data em que a mudança passa a valer (dd/mm/aaaa).'),
                  'E': ('Função', 'Função a partir desta data.'),
                  'F': ('Valor HH', 'Valor-hora a partir desta data.'),
                  'G': ('Turno', 'Turno a partir desta data.'),
                  'H': ('Estilo do contrato', 'Normal / Fixo, Ferista, SPOT ou Compensação a partir desta data.')}
else:
    PROMPT_CAD = {'A': ('RG / matrícula', 'Matrícula da pessoa (igual ao relógio de ponto). É por ela que a macro de pontos acha a pessoa.'),
                  'B': ('Funcionário', 'Nome completo, igual ao usado nos lançamentos.'),
                  'C': ('Situação', 'Ativo ou Desligado. Não apague quem saiu.'),
                  'D': ('Observação', 'Ex.: data de adesão ou de desligamento.')}
    PROMPT_VIG = {'A': ('Funcionário', 'Escolha pela lista. O RG vem sozinho.'),
                  'C': ('Início', 'Data em que a mudança passa a valer (dd/mm/aaaa).'),
                  'D': ('Equipe', 'Serralheria ou SPOT a partir desta data.'),
                  'E': ('Função', 'Função a partir desta data (o Valor HH vem da aba Tabelas).'),
                  'F': ('Turno', 'Turno a partir desta data.'),
                  'G': ('Tipo de contrato', 'Normal / Fixo, Ferista ou SPOT a partir desta data.'),
                  'H': ('Observação', 'Motivo da mudança.')}
COM_CAD = {'A': 'Uma linha por pessoa. Mudou função, turno ou valor? Não edite: ADICIONE uma linha na aba Vigencias. Saiu? Marque Desligado (não apague).'}
COM_VIG = {'A': 'Mudou equipe, função, turno ou tipo de contrato? ADICIONE uma linha com a data de início. Nunca apague nem edite linhas antigas: elas calculam os meses passados.'}


def aplicar_prompts(ws, prompts, r0, r1):
    cols = {CI(c) for c in prompts}
    antigos = list(ws.data_validations.dataValidation)
    params = {}
    ws.data_validations.dataValidation = []
    for dv in antigos:
        resto = []
        for rg in dv.sqref.ranges:
            if rg.min_col == rg.max_col and rg.min_col in cols:
                params.setdefault(rg.min_col, dv)
            else:
                resto.append(rg.coord)
        if resto:
            nv = copy.copy(dv)
            nv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(' '.join(resto))
            ws.data_validations.dataValidation.append(nv)
    for col, (tit, txt) in prompts.items():
        base = params.get(CI(col))
        if base is not None and base.type:
            dv = DataValidation(type=base.type, formula1=base.formula1, formula2=base.formula2, operator=base.operator,
                                allow_blank=True, showErrorMessage=base.showErrorMessage, errorTitle=base.errorTitle,
                                error=base.error, errorStyle=base.errorStyle)
        else:
            dv = DataValidation(allow_blank=True)
        dv.showInputMessage = True
        dv.promptTitle = tit[:32]
        dv.prompt = txt[:255]
        dv.add(f'{col}{r0}:{col}{r1}')
        ws.data_validations.dataValidation.append(dv)


def aplicar_comentarios(ws, coms, hr):
    for col, txt in coms.items():
        c = ws[f'{col}{hr}']
        cm = Comment(txt, 'Controle de Horas TK')
        cm.width, cm.height = 320, 110
        c.comment = cm


aplicar_prompts(wb[L['lanc']], prompts_lanc(), L['l_ini'], L['l_fim'])
aplicar_comentarios(wb[L['lanc']], comentarios_lanc(), L['l_hdr'])
aplicar_prompts(wb['Lanc_HE'], prompts_he(), L['h_ini'], L['h_fim'])
aplicar_comentarios(wb['Lanc_HE'], comentarios_he(), L['h_hdr'])
aplicar_prompts(wb['Compensacoes'], PROMPT_COMP, L['c_ini'], L['c_fim'])
aplicar_comentarios(wb['Compensacoes'], COM_COMP, 3)
aplicar_prompts(wb['Cadastro'], PROMPT_CAD, L['cad_ini'], L['cad_fim'])
aplicar_comentarios(wb['Cadastro'], COM_CAD, L['cad_hdr'])
aplicar_prompts(wb['Vigencias'], PROMPT_VIG, L['vig_ini'], L['vig_fim'])
aplicar_comentarios(wb['Vigencias'], COM_VIG, L['vig_hdr'])

# ======================================================================================
# 2. Visual igual: cabeçalhos auxiliares cinza, zoom, impressão, guias
# ======================================================================================
for nome, hr in ((L['lanc'], L['l_hdr']), ('Lanc_HE', L['h_hdr']), ('Compensacoes', 3), ('Cadastro', L['cad_hdr']),
                 ('Vigencias', L['vig_hdr'])):
    ws = wb[nome]
    for c in ws[hr]:
        if c.value is not None and ('(aux)' in str(c.value) or str(c.value).startswith('vig#')):
            c.fill = fill(AUXF); c.font = Font(name='Arial', size=9, bold=True, color=CINZA)
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BOXH

paineis = {L['painel'], 'Relatorio', L['resumo'], 'Painel_De_Controle', 'SLA', 'Por_Colaborador', 'Como_Usar'}
for ws in wb.worksheets:
    if ws.sheet_state != 'visible':
        continue
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90 if ws.title in paineis else 85
    if ws.title == 'Relatorio':
        ws.oddHeader.left.text = ws.oddHeader.right.text = ''
        ws.oddFooter.center.text = 'Página &P de &N'
        ws.oddFooter.right.text = f'{L["cod"]} — Relatório de ocorrências'
    else:
        ws.oddHeader.left.text = f'&K4d4d4d{L["cod"]} | thyssenkrupp | ' + ws.title.replace('_', ' ')
        ws.oddFooter.right.text = '&K4d4d4dPágina &P de &N'

# topo do painel principal da EI36 igual ao da Gestão da EI32 (título centralizado, faixa da equipe, logo)
if LAY == 'EI36':
    P = wb['Painel']
    P['B1'] = 'EI36 — GESTÃO DE CONTRATO'
    P['B1'].font = Font(name='Arial', size=16, bold=True, color=TK)
    P['B1'].alignment = Alignment(horizontal='center', vertical='center')
    P.merge_cells('B1:I1')
    P['B4'] = 'SERRALHERIA'
    for col in 'BCDEFGHI':
        P[f'{col}4'].fill = fill(TK); P[f'{col}4'].border = BOXH
    P['B4'].font = Font(name='Arial', size=12, bold=True, color='FFFFFFFF')
    P['B4'].alignment = Alignment(horizontal='center', vertical='center')
    P.merge_cells('B4:I4')
    P.row_dimensions[4].height = 18
    P['K2'] = '=HYPERLINK("#\'Como_Usar\'!A1","Como usar ►")'
    P['K2'].fill = fill(NAVY); P['K2'].font = Font(name='Arial', size=9, bold=True, color='FFFFFFFF')
    P['K2'].alignment = Alignment(horizontal='center', vertical='center')
    if LOGO:
        from openpyxl.drawing.image import Image
        img = Image(LOGO)
        img.width, img.height = 64, 50
        P.add_image(img, 'M1')
    # Conferencia / Relatorio já seguem a EI32

# ======================================================================================
# 3. Guia "Como_Usar" — mesmo modelo nas duas planilhas
# ======================================================================================
for velho in ('LEIA-ME', 'Como_Usar'):
    if velho in wb.sheetnames:
        del wb[velho]
G = wb.create_sheet('Como_Usar', 0)
G.sheet_properties.tabColor = 'FF7FC4E8'
G.sheet_view.showGridLines = False
G.sheet_view.zoomScale = 90
for col, w in {'A': 2, 'B': 34, 'C': 44, 'D': 52, 'E': 30, 'F': 2}.items():
    G.column_dimensions[col].width = w
row = [1]


def titulo(txt, link=None):
    r = row[0]
    G[f'B{r}'] = f'=HYPERLINK("#\'{link}\'!A1","{txt}")' if link else txt
    G[f'B{r}'].font = Font(name='Arial', size=16, bold=True, color=TK)
    G.merge_cells(f'B{r}:E{r}'); G.row_dimensions[r].height = 30
    row[0] += 1


def sub(txt):
    r = row[0]
    G[f'B{r}'] = txt
    G[f'B{r}'].font = Font(name='Arial', size=9, italic=True, color=CINZA)
    G[f'B{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    G.merge_cells(f'B{r}:E{r}'); G.row_dimensions[r].height = 28
    row[0] += 2


def faixa(txt):
    r = row[0]
    for col in 'BCDE':
        G[f'{col}{r}'].fill = fill(TK); G[f'{col}{r}'].border = BOXH
    G[f'B{r}'] = txt
    G[f'B{r}'].font = Font(name='Arial', size=10, bold=True, color='FFFFFFFF')
    G[f'B{r}'].alignment = Alignment(vertical='center', indent=1)
    G.merge_cells(f'B{r}:E{r}'); G.row_dimensions[r].height = 20
    row[0] += 1


def cab(*txts):
    r = row[0]
    cols = 'BCDE'
    for i, t in enumerate(txts):
        c = G[f'{cols[i]}{r}']; c.value = t
        c.fill = fill(F2); c.font = Font(name='Arial', size=9, bold=True, color=TXT); c.border = BOX
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    row[0] += 1


def linha(rotulo, texto, link=None, alt=None):
    r = row[0]
    c = G[f'B{r}']
    c.value = f'=HYPERLINK("#\'{link}\'!A1","{rotulo}")' if link else rotulo
    c.font = Font(name='Arial', size=9, bold=True, color=TK if link else TXT, underline='single' if link else None)
    c.alignment = Alignment(vertical='center', wrap_text=True); c.border = BOX
    t = G[f'C{r}']; t.value = texto
    t.font = Font(name='Arial', size=9, color=TXT); t.alignment = Alignment(vertical='center', wrap_text=True)
    for col in 'CDE':
        G[f'{col}{r}'].border = BOX
    G.merge_cells(f'C{r}:E{r}')
    G.row_dimensions[r].height = alt or max(18, 13 * (1 + len(texto) // 120))
    row[0] += 1


def linha4(a, b, c, d, desconta=None):
    r = row[0]
    for col, v in zip('BCDE', (a, b, c, d)):
        x = G[f'{col}{r}']; x.value = v
        x.font = Font(name='Arial', size=9, bold=(col == 'B'), color=TXT)
        x.alignment = Alignment(vertical='center', wrap_text=True); x.border = BOX
    if desconta is not None:
        G[f'E{r}'].fill = fill('FFFADBD8' if desconta else 'FFD5F5E3')
    G.row_dimensions[r].height = max(24, 12 * (1 + max(len(str(v or '')) for v in (b, c)) // 48))
    row[0] += 1


def espaco():
    row[0] += 1


cod, lanc, pnl = L['cod'], L['lanc'], L['painel']
titulo(f'COMO USAR — Lançamento de Horas {cod} ({L["equipes"]})   ◄ {pnl.replace("_", " ")}', pnl)
sub('Guia rápido para o gestor: onde fica cada coisa, como lançar cada situação, o que a planilha faz sozinha e como fechar o mês. '
    'Dica: clique numa célula de lançamento para ver a explicação da coluna; passe o mouse no cabeçalho (triângulo vermelho) para a regra.')

faixa('ONDE CADA COISA FICA')
abas = [(pnl, f'Painel principal. Escolha o mês em {L["mes"]}: contrato, abatimento, valor a pagar, de onde vem cada real descontado.'),
        ('Relatorio', 'Fonte do PDF do mês (botão "Criar Relatório"): pendências, ocorrências por tipo, ferista, detalhe dia a dia, HE e compensações.'),
        (L['resumo'], 'Motor do mês: os 12 meses lado a lado (horas, abatimento, compensação, Valor/Função, SLA, valor final).'),
        (lanc, f'Uma linha por pessoa por dia. Cabeçalho azul = você preenche. A coluna {L["explica"]} explica a linha; a {L["alerta"]} diz o que fazer.'),
        ('Lanc_HE', 'Horas extras: uma linha por pessoa por dia de HE, com o tratamento (obrigatório).'),
        ('Compensacoes', 'Quem trabalhou para compensar, quantas horas e se foi aprovado. Só conta com "Aprovado? = Sim".'),
        ('Conferencia_Ponto', 'Lista montada pelo botão "Atualizar Pontos" com tudo que não bate com o relógio.'),
        ('Cadastro', 'Uma linha por pessoa (RG = matrícula do relógio, nome, Ativo/Desligado).'),
        ('Vigencias', 'Histórico por data de equipe, função, turno e tipo de contrato. Mudou algo: ADICIONE uma linha.'),
        ('Tabelas', f'Regras dos status, funções e valor HH por data, turnos, parâmetros, valor do contrato por data e calendário PCAM.')]
if LAY == 'EI32':
    abas.insert(1, ('Painel_De_Controle', 'Detalhe do mês por equipe e Painel de Verificação do ano (bloco "deve ficar em zero").'))
    abas.insert(4, ('SLA', 'Desconto SLA em R$ por mês e equipe, indicadores e metas.'))
else:
    abas.insert(3, ('Por_Colaborador', 'O mesmo mês do Painel aberto por pessoa e por função.'))
for a, t in abas:
    linha(a, t, link=a)
espaco()

faixa('ROTINA DO DIA')
passos = [('1. Nova linha', f'Em {lanc}: Data, Funcionário (lista) e Status (lista). Um lançamento por pessoa por dia útil.'),
          ('2. Ponto', f'Cole o ponto na coluna {L["ponto"]} (entrada#saída = total) ou use o botão "Atualizar Pontos", que preenche e confere tudo.'),
          ('3. Exceções', 'Atraso, falta, atestado, férias, posto vago, troca de turno: veja a tabela abaixo (COMO LANÇAR CADA SITUAÇÃO).'),
          ('4. Hora extra', 'Lanc_HE com o ponto do dia e o tratamento. Se for compensação: tratamento HORAS COMPENSADAS + uma linha na aba Compensações.'),
          ('5. Confira', f'Coluna {L["alerta"]} (alerta) vazia = linha certa. Coluna {L["explica"]} mostra de onde vem o desconto. Linha vermelha = ponto a conferir.')]
for a, t in passos:
    linha(a, t)
espaco()

faixa('COMO LANÇAR CADA SITUAÇÃO')
cab('Situação', f'Status (coluna {L["status"]})', 'O que mais preencher', 'Desconta do contrato?')
atip = f'{L["atip"]} + horas em {L["atip_desc"]} (desconta) / {L["atip_just"]} (justifica)'
sem_ponto = (f'Sem ponto: horas ausentes em {L["aus_man"]}' if LAY == 'EI36' else 'Sem ponto válido: desconta a jornada inteira')
sit = [
    ('Dia normal', 'Horas Trabalhadas', f'Ponto em {L["ponto"]}.', 'Não. Leve atraso 10–30 min pelo ponto desconta (troque o Status).', False),
    ('Falta sem documento', 'Ausências S/ Justificativa', '—', 'Sim, a jornada (menos o que o ferista ou a compensação cobrir).', True),
    ('Falta com atestado / declaração', 'Ausências C/ Justificativa', f'Observação ({L["obs"]}): só "Dia Atestado".', 'Não.', False),
    ('Atraso ou saída sem documento (inclusive leve, 10–30 min)', 'Entrada Atrasada ou Saída Antecipada S/ Justificativa',
     f'Ponto em {L["ponto"]}. {sem_ponto}.', 'Sim, jornada − ponto.', True),
    ('Atraso ou saída com documento', 'Entrada Atrasada ou Saída Antecipada C/ Justificativa', f'Ponto + motivo na Observação ({L["obs"]}).', 'Não.', False),
    ('Férias', 'Ausências Devido a Férias', '—', 'Sim, se o ferista não cobrir (férias têm prioridade na cobertura).', True),
    ('Posto vago (ninguém contratado)', 'Ausências Devido Falta de Contratação', f'Funcionário VAZIO + função do posto em {L["posto"]}.', 'Sim, a jornada inteira.', True),
    ('Falta de EPI / uniforme', 'Ausências Devido Falta de EPI', 'Motivo na Observação.', 'Sim.', True),
    ('Recesso', 'Ausências Devido a Recesso', '—', 'Sim.', True),
    ('Treinamento', 'Ausências Devido Treinamento', 'Qual treinamento na Observação.', 'Não.', False),
    ('Integração / início de atividades', 'Ausências Devido a Integração / Início de Atividades', 'Observação.', 'Não.', False),
    ('Remanejado para outra atividade', 'Remanejamento', 'Para onde foi, na Observação.', 'Sim, salvo cobertura do ferista.', True),
    ('Afastamento por acidente', 'Ausências Devido a Afastamento por Acidente', 'Motivo na Observação.', 'Não.', False),
    ('Recusou a atividade', 'Recusa de Atividade S/ (ou C/) Motivo Legítimo', 'Motivo na Observação.', 'S/ motivo: Sim. C/ motivo: Não.', True),
    ('Desligamento / pedido de demissão', 'Ausências Devido a Desligamento / Pedido de Demissão', 'Motivo na Observação.', 'Sim (posto descoberto).', True),
    ('Trabalhou para compensar (ex.: sábado)', 'Horas Compensadas', 'Ponto do dia + uma linha na aba Compensações (aprovação).', 'Abate depois de aprovada (em bloco no mês).', False),
    ('Pessoa a mais, não pedida', 'Adicional Não Solicitado', 'Ponto do dia.', 'Abate o absenteísmo do mês 1:1 (não é ferista).', False),
    ('Turno diferente só neste dia', 'O status do dia', f'Turno do dia em {L["turno"]}.', '—', None),
    ('Evento fora do padrão (jogo, evento)', 'O status do dia', f'Situação atípica: {atip}. Motivo na Observação.', 'Horas não acordadas descontam.', True),
    ('Acordo com o cliente', 'O status do dia', f'Horas acordadas em {L["acordo"]} + o acordo na Observação.', 'Não (horas perdoadas).', False),
    ('Hora extra (fim de semana, feriado ou além da jornada)', 'Aba Lanc_HE', 'Horas totais, ponto e Tratamento (obrigatório).', 'Pagas por EI entram no saldo; compensadas abatem faltas.', None),
]
if LAY == 'EI32':
    sit.insert(-1, ('Dívida de horas combinada (ex.: Copa)', 'O status do dia', 'Horas pendentes em Q.', 'Não agora; quitada pelas compensações aprovadas.', False))
for s in sit:
    linha4(*s)
espaco()

faixa('O QUE A PLANILHA FAZ SOZINHA')
regras = [
    ('Jornada', 'ADM: 9h seg–qui e 8h sex. Turnos 2 e 3: seg–sáb. ADM que entra antes das 07:00 e sai até 14:00 vira Turno 2 naquele dia (se sair depois das 14:00, gera alerta).'),
    ('Ferista', 'Cobre só técnicos/serralheiros Normal/Fixo da mesma equipe, nunca o encarregado. Ordem: FÉRIAS primeiro, depois a MAIOR falta. A sobra é doada à próxima falta se for de pelo menos 1h. O ferista nunca gera desconto próprio.'),
    ('Valor/Função', 'Ferista de cargo menor cobrindo cargo maior: desconta (HH coberto − HH ferista) × horas cobertas. Cargo igual ou maior: zero.'),
    ('Compensação', L['comp_regra'] + ' Só conta "Aprovado? = Sim".'),
    ('Leve atraso', 'Atraso de 10 a 30 min pelo ponto desconta. Acima de 30 min com "Horas Trabalhadas" pede para CLASSIFICAR o Status.'),
    ('SLA', f'Desconto manual em R$, informado em {L["sla"]}; sai do valor final a pagar.'),
]
if LAY == 'EI36':
    regras.append(('SPOT', 'Equipe à parte (contrato próprio): aparece em "Horas SPOT" e não entra no abatimento da Serralheria.'))
for a, t in regras:
    linha(a, t)
espaco()

faixa('BOTÕES (suplemento "Controle de Horas TK")')
bt = [('Atualizar Pontos', f'Lê a planilha de Pontos (aba "Base unificada"), preenche o ponto, pinta de vermelho o que não bate com o relógio e cria as linhas de dias com ponto e sem lançamento. Tudo aparece na aba Conferencia_Ponto.'),
      ('Linha vermelha', f'Analise, corrija se precisar e marque "Conferido? = Sim" (coluna {L["ok"]} em {lanc}, {Lh["ok"]} em Lanc_HE).'),
      ('Criar Relatório (PDF)', f'Gera o PDF do mês escolhido em {pnl} {L["mes"]}, em Relatorios\\<nn - {cod}>\\<ano>\\<mm - Mês> - {cod}.pdf. O mesmo mês é substituído.'),
      ('Trocar arquivo de Pontos', f'Escolhe outro arquivo de pontos (fica gravado em {L["param_pontos"]}).'),
      ('Instalar botões', 'Uma vez por planilha (ou se os botões sumirem): Suplementos › Controle de Horas TK › Instalar botões nesta planilha › Ctrl + S.')]
if L['pontos_desde']:
    bt.insert(1, ('A partir de quando', f'O Atualizar Pontos só age a partir da data em {L["pontos_desde"]} (protege meses já fechados).'))
for a, t in bt:
    linha(a, t)
espaco()

faixa('FECHAMENTO DO MÊS')
fech = [('1', f'Alertas do mês resolvidos ({lanc} coluna {L["alerta"]}) e Painel de Verificação em zero ou explicado.'),
        ('2', 'Atualizar Pontos e conferir todas as linhas vermelhas (Conferido? = Sim).'),
        ('3', 'Compensações do mês aprovadas ou recusadas (nenhuma Pendente / PREENCHER).'),
        ('4', 'Toda HE do mês com tratamento e ponto conferido.'),
        ('5', f'Desconto SLA informado ({L["sla"]}), se houver.'),
        ('6', f'{pnl} no mês certo ({L["mes"]}) e "Criar Relatório (PDF)". Conferir o PDF com o painel ("✔ confere").')]
for a, t in fech:
    linha(a, t)
espaco()

faixa('MUDANÇAS DE CADASTRO')
mud = [('Pessoa nova', 'Cadastro: RG (matrícula do relógio) + nome + Ativo. Vigencias: uma linha com a data de início, equipe, função, turno e tipo de contrato.'),
       ('Mudou função / turno / virou ferista', 'Vigencias: ADICIONE uma linha com a data da mudança. Não edite as linhas antigas.'),
       ('Reajuste de valor HH ou do contrato', f'Aba Tabelas: ADICIONE uma linha com a nova data de início e o novo valor ({L["contrato"]} para o contrato).'),
       ('Desligamento', 'Cadastro: marque Desligado. Não apague a pessoa (os meses passados dependem dela).')]
for a, t in mud:
    linha(a, t)
espaco()

faixa('CORES E AVISOS')
cores = [('Cabeçalho azul / célula cinza-claro', 'Você preenche.'),
         ('Cabeçalho marinho', 'Automático (fórmula). Não digite por cima.'),
         ('Cabeçalho cinza', 'Colunas auxiliares, ficam recolhidas ([+] acima da coluna).'),
         ('Vermelho-claro', 'Precisa de atenção: ponto divergente, desconto, alerta.'),
         ('Verde-claro', 'Conferido / reduziu o desconto (ferista, compensação, acordo).'),
         ('Triângulo vermelho no cabeçalho', 'Passe o mouse: explica a regra da coluna.'),
         ('Mensagem ao clicar na célula', 'Explica o que digitar e o formato (h:mm, dd/mm/aaaa, lista).')]
for a, t in cores:
    linha(a, t)
G.print_area = f'B1:E{row[0]}'
G.page_setup.orientation = 'portrait'
G.page_setup.fitToWidth = 1; G.page_setup.fitToHeight = 0
G.sheet_properties.pageSetUpPr.fitToPage = True
G.oddHeader.left.text = f'&K4d4d4d{cod} | thyssenkrupp | Como usar'
G.oddFooter.right.text = '&K4d4d4dPágina &P de &N'

# abre no painel principal
wb.active = wb.sheetnames.index(pnl)
for ws in wb.worksheets:
    ws.sheet_view.tabSelected = (ws.title == pnl)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('ok', OUT)
