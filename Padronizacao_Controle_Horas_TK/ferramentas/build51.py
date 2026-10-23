"""v51 — relatório completo: toda situação fora da rotina, situação por situação, + HE por tratamento."""
import openpyxl, copy, sys
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import column_index_from_string as CI
from openpyxl.worksheet.pagebreak import Break, RowBreak

SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)
lh = wb['Lanc_HORAS']; T = wb['Tabelas']; ws = wb['Relatorio']
TK = 'FF0095D8'; NAVY = 'FF1F3864'; TXT = 'FF4D4D4D'; CINZA = 'FF7F7F7F'
VERDE, VERM = 'FFD5F5E3', 'FFFADBD8'
fill = lambda c: PatternFill('solid', fgColor=c)
ln = Side(style='thin', color='FFD9D9D9'); B = Border(left=ln, right=ln, top=ln, bottom=ln)
def A(h='center', wrap=False, ind=0): return Alignment(horizontal=h, vertical='center', wrap_text=wrap, indent=ind)
G = 'Gestao_Contrato!$B$3'
NL = 3884

# =====================================================================
# 1. Tabelas: rótulos dos status que faltavam
# =====================================================================
rot = {T.cell(r, 28).value: r for r in range(3, 41) if T.cell(r, 28).value}
prox = max(rot.values()) + 1
for st, lab in [('Horas Extras', 'HE lançada na Lanc_HORAS'), ('Situação Atípica', 'Situação atípica'),
                ('Recusa de Atividade S/ Motivo Legítimo', 'Recusa de atividade (sem motivo legítimo)'),
                ('Recusa de Atividade C/ Motivo Legítimo', 'Recusa de atividade (motivo legítimo)'),
                ('Ausências Devido a Integração / Início de Atividades', 'Integração / início de atividades'),
                ('Ausências Devido a Afastamento por Acidente', 'Afastamento por acidente'),
                ('Ausências Devido a Desligamento / Pedido de Demissão', 'Desligamento / pedido de demissão'),
                ('Adicional Não Solicitado', 'Adicional não solicitado')]:
    if st not in rot:
        for c, v in ((28, st), (29, lab)):
            x = T.cell(prox, c); x.value = v; x._style = copy.copy(T.cell(prox - 1, c)._style)
        prox += 1

# =====================================================================
# 2. Lanc_HORAS: AQ (situação) + AX (sem motivo)
# =====================================================================
MOTIVO_OBRIG = ['Falta justificada (atestado/comprovante)', 'Atraso/saída justificado (comprovante)', 'Remanejamento',
                'Situação atípica', 'Acordo', 'Ajuste manual', 'Treinamento',
                'Recusa de atividade (sem motivo legítimo)', 'Recusa de atividade (motivo legítimo)',
                'Afastamento por acidente', 'Desligamento / pedido de demissão']
lh['AX1'] = 'Sem motivo (aux)'; lh['AX1']._style = copy.copy(lh['AU1']._style)
for r in range(2, NL + 1):
    turno = f'AND(TRIM($S{r})<>"",UPPER(TRIM($S{r}))<>UPPER(TRIM($H{r})))'
    pend = f'AND($AV{r}<>"",$AW{r}<>"Sim")'
    lh[f'AQ{r}'] = (
        f'=IF($A{r}="","",IF(AND($C{r}="Ferista",$J{r}="Horas Trabalhadas"),IF({pend},"Ponto divergente (a conferir)",""),'
        f'IF($L{r}<>"","Situação atípica",'
        f'IF($J{r}="Horas Trabalhadas",IF(N($P{r})>0,"Acordo",IF(N($Q{r})>0,"Pendência de horas",IF(N($V{r})>0,"Leve atraso (10 a 30 min)",'
        f'IF(N($M{r})+N($N{r})>0,"Ajuste manual",IF({turno},"Troca de turno",IF({pend},"Ponto divergente (a conferir)","")))))),'
        f'IF(AND($J{r}="Entrada Atrasada ou Saída Antecipada S/ Justificativa",N($V{r})>0,N($V{r})<=30/1440),"Leve atraso (10 a 30 min)",'
        f'IFERROR(INDEX(Tabelas!$AC$3:$AC$40,MATCH($J{r},Tabelas!$AB$3:$AB$40,0)),$J{r}))))))')
    ors = ','.join(f'$AQ{r}="{m}"' for m in MOTIVO_OBRIG)
    lh[f'AX{r}'] = f'=IF(AND($AQ{r}<>"",TRIM($X{r})="",OR({ors})),1,0)'
    lh[f'AX{r}']._style = copy.copy(lh[f'AU{r}']._style)
lh.column_dimensions['AX'].hidden = True

# =====================================================================
# 3. Relatório: limpa da linha 9 em diante e reconstrói
# =====================================================================
OLD_LAST = int(ws['L1'].value)
for m in list(ws.merged_cells.ranges):
    if m.min_row >= 9: ws.unmerge_cells(str(m))
for r in range(9, OLD_LAST + 5):
    for c in range(1, 15):
        x = ws.cell(r, c); x.value = None; x.style = 'Normal'
    if r in ws.row_dimensions: ws.row_dimensions[r].height = None
ws.conditional_formatting = type(ws.conditional_formatting)()
ws.row_breaks = RowBreak()
quebra = lambda rr: ws.row_breaks.append(Break(id=rr - 1))
for k_, w_ in {'B': 11, 'C': 28, 'D': 26, 'E': 8, 'F': 17, 'G': 9, 'H': 11, 'I': 52, 'J': 26}.items():
    ws.column_dimensions[k_].width = w_

def cell(rc, v=None, size=9, b=False, color=TXT, fl=None, h='left', wrap=False, fmt=None, brd=True, ind=0):
    x = ws[rc]
    if v is not None: x.value = v
    x.font = Font(name='Arial', size=size, bold=b, color=color); x.alignment = A(h, wrap, ind)
    if fl: x.fill = fill(fl)
    if brd: x.border = B
    if fmt: x.number_format = fmt
    return x
def bar(r, text, color=TK, size=10, hide=None):
    ws.merge_cells(f'B{r}:J{r}'); cell(f'B{r}', text, size=size, b=True, color='FFFFFFFF', fl=color, ind=1, brd=False)
    for col in range(2, 11): ws.cell(r, col).fill = fill(color)
    ws.row_dimensions[r].height = 20
    if hide: ws[f'N{r}'] = hide
def sub(r, text, hide=None):
    ws.merge_cells(f'B{r}:J{r}'); cell(f'B{r}', text, size=9, b=True, color=NAVY, fl='FFDDEBF7', ind=1, brd=False)
    for col in range(2, 11): ws.cell(r, col).fill = fill('FFDDEBF7')
    ws.row_dimensions[r].height = 16
    if hide: ws[f'N{r}'] = hide
def hdrrow(r, cols, hide=None):
    for c, t in cols: cell(f'{c}{r}', t, b=True, fl='FFF2F2F2', h='center', wrap=True)
    ws.row_dimensions[r].height = 26
    if hide: ws[f'N{r}'] = hide
def note(r, text, h=16):
    ws.merge_cells(f'B{r}:J{r}'); cell(f'B{r}', text, size=8, color=CINZA, brd=False, wrap=True); ws.row_dimensions[r].height = h

ws['B3'].value = ('Tudo o que não foi um dia normal de trabalho: faltas, atrasos (leves ou não), férias, atestados e comprovantes, '
                  'remanejamentos, trocas de turno, cobertura do ferista (cobriu ou não), acordos, compensações, divergências de ponto '
                  'e todas as horas extras com o tratamento de cada uma. O mês segue o filtro da Gestão de Contrato (B3).')
ws.row_dimensions[3].height = 30

LHr = lambda c: f'Lanc_HORAS!${c}$2:${c}${NL}'
HEr = lambda c: f'Lanc_HE!${c}$2:${c}$693'
CPr = lambda c: f'Compensacoes!${c}$4:${c}$503'
r = 10

# ---------- PENDÊNCIAS ----------
bar(r, 'PENDÊNCIAS PARA FECHAR O MÊS (tudo precisa estar em zero)', color=NAVY); r += 1
hdrrow(r, [('B', 'Qtde'), ('C', 'Pendência'), ('D', ''), ('E', ''), ('F', ''), ('G', ''), ('H', ''), ('I', 'O que fazer'), ('J', '')])
ws.merge_cells(f'C{r}:H{r}'); ws.merge_cells(f'I{r}:J{r}'); r += 1
pends = [
    (f'=SUMPRODUCT(({LHr("B")}={G})*({LHr("AV")}<>"")*({LHr("AW")}<>"Sim"))',
     'Lanc_HORAS: linhas com ponto divergente ainda não conferidas', 'Analisar a linha vermelha e marcar "Conferido? = Sim" (coluna AW).'),
    (f'=SUMPRODUCT(({HEr("B")}={G})*({HEr("AH")}<>"")*({HEr("AI")}<>"Sim"))',
     'Lanc_HE: horas extras com ponto divergente não conferidas', 'Conferir o ponto e marcar "Conferido? = Sim" (coluna AI).'),
    (f'=SUMPRODUCT(({HEr("B")}={G})*({HEr("A")}<>"")*({HEr("V")}=""))',
     'Horas extras sem tratamento definido', 'Escolher o tratamento (pagas por EI, pela área, compensadas ou remanejadas).'),
    (f'=SUMPRODUCT(({CPr("H")}={G})*({CPr("A")}<>"")*({CPr("F")}<>"Sim"))',
     'Compensações não aprovadas (pendentes ou recusadas)', 'Só contam com "Aprovado? = Sim".'),
    (f'=SUMPRODUCT(({LHr("B")}={G})*({LHr("AX")}=1))',
     'Ocorrências sem motivo na Observação (justificadas, remanejamento, atípicas, acordos, ajustes, treinamento)',
     'Preencher a coluna Observação (X) da Lanc_HORAS.'),
]
p0 = r
for f, t, o in pends:
    cell(f'B{r}', f, h='center', b=True)
    ws.merge_cells(f'C{r}:H{r}'); cell(f'C{r}', t)
    for c in 'DEFGH': ws[f'{c}{r}'].border = B
    ws.merge_cells(f'I{r}:J{r}'); cell(f'I{r}', o, size=8, color=CINZA, wrap=True); ws[f'J{r}'].border = B
    ws.row_dimensions[r].height = 16; r += 1
ws.conditional_formatting.add(f'B{p0}:B{r-1}', FormulaRule(formula=[f'N(B{p0})=0'], fill=fill(VERDE)))
ws.conditional_formatting.add(f'B{p0}:B{r-1}', FormulaRule(formula=[f'N(B{p0})>0'], fill=fill(VERM)))
r += 1

# ---------- FERISTA ----------
bar(r, 'FERISTA DO MÊS'); r += 1
hdrrow(r, [('B', 'Equipe'), ('C', 'Ferista'), ('D', 'Trabalhou (h)'), ('E', 'Cobert. (qtde)'), ('F', 'Cobriu (h)'),
           ('G', 'Sem cobertura (h)'), ('H', 'Faltas injust. do ferista (h)'), ('I', 'Leitura'), ('J', '')]); ws.merge_cells(f'I{r}:J{r}'); r += 1
for team, ar in (('Elétrica', 95), ('Climatização', 96)):
    cell(f'B{r}', team, b=True); cell(f'C{r}', f'=IF(Analises!$N${ar}="","—",Analises!$N${ar})')
    cell(f'D{r}', f'=N(Analises!$O${ar})', h='center', fmt='[h]:mm')
    cell(f'E{r}', f'=N(Analises!$R${ar})', h='center')
    cell(f'F{r}', f'=N(Analises!$P${ar})', h='center', fmt='[h]:mm', b=True, color=TK)
    cell(f'G{r}', f'=N(Analises!$Q${ar})', h='center', fmt='[h]:mm')
    cell(f'H{r}', f'=N(Analises!$S${ar})', h='center', fmt='[h]:mm')
    ws.merge_cells(f'I{r}:J{r}')
    cell(f'I{r}', 'Cobriu = horas de falta/atraso/férias da equipe que o ferista assumiu (não abatem). '
                  'Sem cobertura = horas que ele trabalhou sem falta para cobrir.', size=8, color=CINZA, wrap=True)
    ws[f'J{r}'].border = B; ws.row_dimensions[r].height = 24; r += 1
note(r, 'A coluna "Ferista" do detalhe diz, ocorrência por ocorrência, se o ferista cobriu tudo, parte ou nada.'); r += 2

# ---------- OCORRÊNCIAS POR TIPO + DETALHE ----------
grupos = [
    ('DESCONTAM DO CONTRATO (salvo o que o ferista, acordo ou compensação cobrir)', [
        ('Falta sem justificativa', 'Dia inteiro sem documento. Desconta a jornada, menos o que o ferista ou uma compensação cobriu.'),
        ('Atraso/saída sem justificativa', 'Atraso ou saída antecipada acima de 30 min, sem documento. Desconta o tempo fora.'),
        ('Leve atraso (10 a 30 min)', 'Atraso ou saída de 10 a 30 min sem documento. Desconta o tempo; o ferista cobre se sobrar hora no dia.'),
        ('Posto vago (falta de contratação)', 'Vaga sem ninguém contratado. Desconta a jornada inteira.'),
        ('Remanejamento', 'Pessoa deslocada para outra atividade/área. O posto fica descoberto e desconta, salvo cobertura do ferista.'),
        ('Férias', 'Desconta a jornada, menos o que o ferista ou acordo cobriu.'),
        ('Recesso', 'Ausência por recesso. Conta no absenteísmo (Tabelas).'),
        ('Falta de EPI', 'Sem EPI para trabalhar. Conta no absenteísmo.'),
        ('Situação atípica', 'Evento fora da rotina (ex.: jogo, evento da empresa). Desconta o que não foi acordado.'),
        ('Recusa de atividade (sem motivo legítimo)', 'Presente, mas se recusou a executar o serviço sem motivo legítimo. Desconta a jornada do dia.'),
        ('Desligamento / pedido de demissão', 'Dia do desligamento ou do pedido de demissão. Desconta; depois disso o posto vira posto vago.'),
    ]),
    ('NÃO DESCONTAM (justificadas)', [
        ('Falta justificada (atestado/comprovante)', 'Ausência com atestado/declaração; Observação "Dia Atestado" (sem o motivo médico).'),
        ('Atraso/saída justificado (comprovante)', 'Atraso ou saída com declaração de comparecimento.'),
        ('Treinamento', 'Ausência para treinamento.'),
        ('Recusa de atividade (motivo legítimo)', 'Recusou a atividade por motivo legítimo (ex.: risco à segurança). Não desconta; o motivo é obrigatório.'),
        ('Integração / início de atividades', 'Dia de integração TK ou início antes de assumir o posto. Não desconta.'),
        ('Afastamento por acidente', 'Afastado por acidente ou lesão. Não desconta; registrar o motivo/CAT na Observação.'),
    ]),
    ('AJUSTES, ACORDOS E CONTROLES', [
        ('Troca de turno', 'Trabalhou em turno diferente do cadastro (ex.: ADM → TURNO 2). Não desconta; a jornada do dia segue o turno trabalhado.'),
        ('Acordo', 'Horas perdoadas por acordo comercial. Não descontam.'),
        ('Pendência de horas', 'Horas devidas; quitadas depois por compensação aprovada.'),
        ('Ajuste manual', 'Horas lançadas à mão (colunas M/N). Conferir o motivo.'),
        ('Compensação (dia trabalhado)', 'Dia trabalhado para compensar horas (ver Compensações do mês).'),
        ('Reposição de férias', 'Controle de reposição de férias. Não desconta.'),
        ('Banco de horas', 'Controle de banco de horas. Não desconta.'),
        ('Horas transferidas', 'Horas de cobertura passadas para outra pessoa. Não desconta.'),
        ('Horas recebidas', 'Horas de cobertura recebidas de outra pessoa. Não desconta.'),
        ('Adicional não solicitado', 'Pessoa a mais, não pedida. Horas = horas trabalhadas por ela; abatem o absenteísmo do mês 1:1 (como compensação).'),
    ]),
    ('PARA CONFERIR', [
        ('Ponto divergente (a conferir)', 'Dia normal, mas o relógio não bate com o lançamento e ainda não está "Conferido = Sim".'),
        ('HE lançada na Lanc_HORAS', 'Status "Horas Extras" usado na Lanc_HORAS. HE deve ir para a Lanc_HE.'),
    ]),
]
NDET = 150
for team, flag, cum in (('Elétrica', 'AR', 'AS'), ('Climatização', 'AT', 'AU')):
    r += 1
    if team == 'Climatização': quebra(r)
    bar(r, f'{team.upper()} — OCORRÊNCIAS POR TIPO'); r += 1
    hdrrow(r, [('B', 'Qtde'), ('C', 'Situação'), ('D', ''), ('E', 'Horas'), ('F', 'Ferista cobriu (h)'), ('G', 'Desconta (h)'),
               ('H', 'Valor (R$)'), ('I', 'O que significa'), ('J', '')])
    ws.merge_cells(f'C{r}:D{r}'); ws.merge_cells(f'I{r}:J{r}'); r += 1
    linhas_cat = []
    for gnome, cats in grupos:
        g0 = r + 1
        sub(r, gnome, hide=f'=IF(SUM(B{g0}:B{g0+len(cats)-1})=0,1,"")'); r += 1
        for ct, mean in cats:
            cr = f'Lanc_HORAS!$AQ$2:$AQ${NL},"{ct}",Lanc_HORAS!$B$2:$B${NL},{G},Lanc_HORAS!$D$2:$D${NL},"{team}"'
            cell(f'B{r}', f'=COUNTIFS({cr})', h='center', b=True)
            ws.merge_cells(f'C{r}:D{r}'); cell(f'C{r}', ct); ws[f'D{r}'].border = B
            cell(f'E{r}', (f'=SUMIFS({LHr("U")},{cr})' if ct == 'Adicional não solicitado' else f'=SUMIFS({LHr("V")},{cr})+SUMIFS({LHr("M")},{cr})'), h='center', fmt='[h]:mm')
            cell(f'F{r}', f'=SUMIFS({LHr("AC")},{cr})', h='center', fmt='[h]:mm')
            cell(f'G{r}', f'=SUMIFS({LHr("AM")},{cr})', h='center', fmt='[h]:mm')
            cell(f'H{r}', f'=SUMIFS({LHr("W")},{cr})', h='right', fmt='#,##0.00')
            ws.merge_cells(f'I{r}:J{r}'); cell(f'I{r}', mean, size=8, color=CINZA, wrap=True); ws[f'J{r}'].border = B
            ws[f'N{r}'] = f'=IF(B{r}=0,1,"")'; ws.row_dimensions[r].height = 16
            linhas_cat.append(r); r += 1
    # Outros: garante que o total bate com o nº de ocorrências
    soma = '+'.join(f'B{x}' for x in linhas_cat)
    cell(f'B{r}', f'=SUM({LHr(flag)})-({soma})', h='center', b=True)
    ws.merge_cells(f'C{r}:D{r}'); cell(f'C{r}', 'Outros (status sem rótulo)'); ws[f'D{r}'].border = B
    for c in 'EFGH': cell(f'{c}{r}', '', h='center')
    ws.merge_cells(f'I{r}:J{r}'); cell(f'I{r}', 'Status que não tem rótulo em Tabelas › Rótulos do relatório. Deve ficar em zero.', size=8, color=CINZA, wrap=True)
    ws[f'J{r}'].border = B; ws[f'N{r}'] = f'=IF(B{r}=0,1,"")'; linhas_cat.append(r); r += 1
    cell(f'C{r}', 'TOTAL', b=True, fl='FFD9D9D9'); ws.merge_cells(f'C{r}:D{r}')
    for c, fmt in (('B', None), ('E', '[h]:mm'), ('F', '[h]:mm'), ('G', '[h]:mm'), ('H', '#,##0.00')):
        cell(f'{c}{r}', '=' + '+'.join(f'N({c}{x})' for x in linhas_cat), b=True, fl='FFD9D9D9', h='center' if c != 'H' else 'right', fmt=fmt)
    for c in 'DIJ': ws[f'{c}{r}'].fill = fill('FFD9D9D9')
    r += 1
    bs = 3 if team == 'Elétrica' else 19
    IX = lambda col: f'IFERROR(INDEX(Analises!${col}${bs}:${col}${bs+11},MATCH({G},Analises!$A${bs}:$A${bs+11},0)),0)'
    ws.merge_cells(f'C{r}:D{r}'); cell(f'C{r}', '( − ) Compensações aprovadas + adicional não solicitado')
    ws[f'D{r}'].border = B
    cell(f'B{r}', '', h='center'); cell(f'E{r}', '', h='center'); cell(f'F{r}', '', h='center')
    cell(f'G{r}', f'={IX("AW")}+{IX("BG")}', h='center', fmt='[h]:mm')
    cell(f'H{r}', f'={IX("AT")}', h='right', fmt='#,##0.00')
    ws.merge_cells(f'I{r}:J{r}'); cell(f'I{r}', 'Horas que abatem o absenteísmo do mês 1:1 (aba Compensações e status Adicional não solicitado).', size=8, color=CINZA, wrap=True)
    ws[f'J{r}'].border = B; ws[f'N{r}'] = f'=IF(N(G{r})=0,1,"")'; r += 1
    ws.merge_cells(f'C{r}:D{r}'); cell(f'C{r}', 'ABATIMENTO DO MÊS (resultado)', b=True, fl='FFDDEBF7')
    for c in 'BDEFIJ': ws[f'{c}{r}'].fill = fill('FFDDEBF7'); ws[f'{c}{r}'].border = B
    cell(f'G{r}', f'={IX("B")}', b=True, fl='FFDDEBF7', h='center', fmt='[h]:mm')
    cell(f'H{r}', f'={IX("I")}', b=True, fl='FFDDEBF7', h='right', fmt='#,##0.00', color=TK)
    ws.merge_cells(f'I{r}:J{r}')
    r += 2
    if team == 'Elétrica': quebra(r)
    bar(r, f'{team.upper()} — DETALHE DIA A DIA', hide=f'=IF(SUM({LHr(flag)})=0,1,"")'); r += 1
    hdrrow(r, [('B', 'Data'), ('C', 'Funcionário'), ('D', 'Situação'), ('E', 'Horas'), ('F', 'Ferista'), ('G', 'Desconta (h)'),
               ('H', 'Valor (R$)'), ('I', 'Explicação (como o cálculo tratou)'), ('J', 'Motivo / observação')],
           hide=f'=IF(SUM({LHr(flag)})=0,1,"")'); r += 1
    d0 = r
    for k in range(1, NDET + 1):
        ws[f'M{r}'] = k
        ws[f'L{r}'] = f'=IFERROR(MATCH(M{r},Lanc_HORAS!${cum}$2:${cum}${NL},0)+1,"")'
        I = lambda c: f'INDEX(Lanc_HORAS!${c}$1:${c}${NL},$L{r})'
        base = f'(N({I("V")})+N({I("M")}))'
        cov = f'N({I("AC")})'
        turno = f'AND(TRIM({I("S")})<>"",UPPER(TRIM({I("S")}))<>UPPER(TRIM({I("H")})))'
        pend = f'AND({I("AV")}<>"",{I("AW")}<>"Sim")'
        cell(f'B{r}', f'=IF($L{r}="","",{I("A")})', h='center', fmt='dd/mm/yyyy')
        cell(f'C{r}', f'=IF($L{r}="","",IF(TRIM({I("E")})="","Posto vago — "&{I("R")},{I("E")}&IF({I("C")}="Ferista"," (ferista)","")))', wrap=True)
        cell(f'D{r}', f'=IF($L{r}="","",{I("AQ")}&IF(AND({I("AQ")}<>"Troca de turno",{turno}),'
                      f'" · turno "&{I("S")}&" (cadastro "&{I("H")}&")",IF({I("AQ")}="Troca de turno",": "&{I("S")}&" (cadastro "&{I("H")}&")",""))'
                      f'&IF(AND({I("AQ")}<>"Ponto divergente (a conferir)",{pend})," · ponto a conferir",""))', size=8, wrap=True)
        cell(f'E{r}', f'=IF($L{r}="","",IF({I("J")}="Adicional Não Solicitado",N({I("U")}),N({I("V")})+N({I("M")})+N({I("N")})))', h='center', fmt='[h]:mm;;;')
        cell(f'F{r}', f'=IF($L{r}="","",IF({base}=0,"",IF({cov}>={base}-1/1440,"Cobriu tudo ("&TEXT({cov},"[h]:mm")&")",'
                      f'IF({cov}>0,"Parcial: "&TEXT({cov},"[h]:mm")&" de "&TEXT({base},"[h]:mm")&IF(N({I("P")})>0," + acordo",""),'
                      f'IF({I("T")}="Não",IF(N({I("P")})>0,"Não precisou (acordo)","Não desconta"),"Não cobriu")))))', size=8, h='center', wrap=True)
        cell(f'G{r}', f'=IF($L{r}="","",N({I("AM")}))', h='center', fmt='[h]:mm;;;')
        cell(f'H{r}', f'=IF($L{r}="","",N({I("W")}))', h='right', fmt='#,##0.00;;;')
        cell(f'I{r}', f'=IF($L{r}="","",{I("AA")})', size=8, wrap=True)
        cell(f'J{r}', f'=IF($L{r}="","",IF(TRIM({I("X")})="",IF({I("AX")}=1,"⚠ SEM MOTIVO — preencher a Observação",""),{I("X")}))', size=8, wrap=True)
        ws[f'N{r}'] = f'=IF($L{r}="",1,"")'
        if k % 2 == 0:
            for c in 'BCDEFGHIJ': ws[f'{c}{r}'].fill = fill('FFF7F7F7')
        ws.row_dimensions[r].height = 24; r += 1
    ws.conditional_formatting.add(f'J{d0}:J{r-1}', FormulaRule(formula=[f'ISNUMBER(SEARCH("SEM MOTIVO",$J{d0}))'], fill=fill(VERM)))
    ws.conditional_formatting.add(f'F{d0}:F{r-1}', FormulaRule(formula=[f'$F{d0}="Não cobriu"'], fill=fill(VERM)))
    ws.conditional_formatting.add(f'F{d0}:F{r-1}', FormulaRule(formula=[f'LEFT($F{d0},6)="Cobriu"'], fill=fill(VERDE)))
    ws.conditional_formatting.add(f'D{d0}:D{r-1}', FormulaRule(formula=[f'ISNUMBER(SEARCH("ponto",$D{d0}))'], fill=fill(VERM)))
    ws[f'B{r}'] = f'=IF(SUM({LHr(flag)})>{NDET},"ATENÇÃO: o mês tem mais de {NDET} ocorrências; aumente o relatório.","")'
    ws[f'B{r}'].font = Font(name='Arial', size=9, bold=True, color=TXT)
    ws[f'N{r}'] = f'=IF(B{r}="",1,"")'; r += 1

# ---------- HORAS EXTRAS ----------
NHE = 60
hidHE = f'=IF(SUM(Lanc_HE!$AF$2:$AF$693)=0,1,"")'
r += 1; quebra(r); bar(r, 'HORAS EXTRAS DO MÊS — POR TRATAMENTO (efeito no contrato)'); r += 1
hdrrow(r, [('B', 'Qtde'), ('C', 'Tratamento'), ('D', ''), ('E', 'Elétrica (h)'), ('F', 'Climatização (h)'), ('G', 'HE 100% (h)'),
           ('H', 'Valor (R$)'), ('I', 'Efeito no contrato'), ('J', '')]); ws.merge_cells(f'C{r}:D{r}'); ws.merge_cells(f'I{r}:J{r}'); r += 1
trats = [('PAGAS POR EI', 'Paga pela EI à parte do contrato (valor a mais). Também entra no saldo do mês (abatimento − HE paga por EI).'),
         ('PAGAS PELA ÁREA', 'Paga pela área que pediu. Não entra no valor da EI.'),
         ('HORAS COMPENSADAS', 'Não é paga. As horas vão para Compensações e abatem faltas/pendências quando aprovadas.'),
         ('REMANEJADO P/ OUTRA ÁREA', 'Horas feitas para outra área. Não pagas por este contrato.')]
h0 = r
hc = lambda col, crit, extra='': f'SUMIFS({HEr(col)},{HEr("B")},{G},{HEr("V")},"{crit}"{extra})'
for tr, ef in trats:
    cell(f'B{r}', f'=COUNTIFS({HEr("B")},{G},{HEr("V")},"{tr}",{HEr("A")},"<>")', h='center', b=True)
    ws.merge_cells(f'C{r}:D{r}'); cell(f'C{r}', tr, b=(tr == 'PAGAS POR EI')); ws[f'D{r}'].border = B
    cell(f'E{r}', '=' + hc('R', tr, f',{HEr("D")},"Elétrica"'), h='center', fmt='[h]:mm')
    cell(f'F{r}', '=' + hc('R', tr, f',{HEr("D")},"Climatização"'), h='center', fmt='[h]:mm')
    cell(f'G{r}', '=' + hc('T', tr), h='center', fmt='[h]:mm')
    cell(f'H{r}', '=' + hc('U', tr), h='right', fmt='#,##0.00', b=(tr == 'PAGAS POR EI'), color=TK if tr == 'PAGAS POR EI' else TXT)
    ws.merge_cells(f'I{r}:J{r}'); cell(f'I{r}', ef, size=8, color=CINZA, wrap=True); ws[f'J{r}'].border = B
    ws.row_dimensions[r].height = 24; r += 1
# sem tratamento = total − tratados
tot = lambda col, extra='': f'SUMIFS({HEr(col)},{HEr("B")},{G}{extra})'
cell(f'B{r}', f'=COUNTIFS({HEr("B")},{G},{HEr("A")},"<>")-SUM(B{h0}:B{r-1})', h='center', b=True)
ws.merge_cells(f'C{r}:D{r}'); cell(f'C{r}', 'SEM TRATAMENTO', b=True); ws[f'D{r}'].border = B
XE = ',' + HEr('D') + ',"Elétrica"'; XC = ',' + HEr('D') + ',"Climatização"'
cell(f'E{r}', '=' + tot('R', XE) + f'-SUM(E{h0}:E{r-1})', h='center', fmt='[h]:mm')
cell(f'F{r}', '=' + tot('R', XC) + f'-SUM(F{h0}:F{r-1})', h='center', fmt='[h]:mm')
cell(f'G{r}', f'={tot("T")}-SUM(G{h0}:G{r-1})', h='center', fmt='[h]:mm')
cell(f'H{r}', f'={tot("U")}-SUM(H{h0}:H{r-1})', h='right', fmt='#,##0.00')
ws.merge_cells(f'I{r}:J{r}'); cell(f'I{r}', 'Pendente: sem tratamento a HE não entra em nenhum valor. Definir antes de fechar o mês.', size=8, color=CINZA, wrap=True)
ws[f'J{r}'].border = B; ws.row_dimensions[r].height = 24
ws.conditional_formatting.add(f'B{r}', FormulaRule(formula=[f'N(B{r})>0'], fill=fill(VERM)))
r += 1
cell(f'C{r}', 'TOTAL', b=True, fl='FFD9D9D9'); ws.merge_cells(f'C{r}:D{r}')
for c, fmt in (('B', None), ('E', '[h]:mm'), ('F', '[h]:mm'), ('G', '[h]:mm'), ('H', '#,##0.00')):
    cell(f'{c}{r}', f'=SUM({c}{h0}:{c}{r-1})', b=True, fl='FFD9D9D9', h='center' if c != 'H' else 'right', fmt=fmt)
for c in 'DIJ': ws[f'{c}{r}'].fill = fill('FFD9D9D9')
r += 1
note(r, 'Horas = horas líquidas de HE (ponto − almoço). HE 100% = domingo, feriado e o que passar da jornada de sábado; o resto é 50%. '
        'Valor calculado com o valor HE da função; para tratamentos não pagos é só referência.', h=22); r += 2

bar(r, 'HORAS EXTRAS DO MÊS — CASO A CASO', hide=hidHE); r += 1
hdrrow(r, [('B', 'Data'), ('C', 'Funcionário'), ('D', 'Tratamento'), ('E', 'Tipo do dia'), ('F', 'HE 50% (h)'), ('G', 'HE 100% (h)'),
           ('H', 'Valor (R$)'), ('I', 'Descrição · solicitante · área · requisição · pendências'), ('J', 'Equipe')], hide=hidHE); r += 1
e0 = r
for k in range(1, NHE + 1):
    ws[f'M{r}'] = k; ws[f'L{r}'] = f'=IFERROR(MATCH(M{r},Lanc_HE!$AG$2:$AG$693,0)+1,"")'
    I = lambda c: f'INDEX(Lanc_HE!${c}$1:${c}$693,$L{r})'
    cell(f'B{r}', f'=IF($L{r}="","",{I("A")})', h='center', fmt='dd/mm/yyyy')
    cell(f'C{r}', f'=IF($L{r}="","",{I("E")})', wrap=True)
    cell(f'D{r}', f'=IF($L{r}="","",IF({I("V")}="","SEM TRATAMENTO",{I("V")}))', size=8, wrap=True)
    cell(f'E{r}', f'=IF($L{r}="","",{I("O")})', h='center')
    cell(f'F{r}', f'=IF($L{r}="","",N({I("S")}))', h='center', fmt='[h]:mm;;;')
    cell(f'G{r}', f'=IF($L{r}="","",N({I("T")}))', h='center', fmt='[h]:mm;;;')
    cell(f'H{r}', f'=IF($L{r}="","",N({I("U")}))', h='right', fmt='#,##0.00;;;')
    tudo = (f'IF({I("V")}=""," · ⚠ SEM TRATAMENTO","")'
            f'&IF(AND({I("AH")}<>"",{I("AI")}<>"Sim")," · ⚠ PONTO A CONFERIR","")'
            f'&IF({I("AA")}="",""," · "&{I("AA")})'
            f'&IF({I("Y")}="",""," · Solicitante: "&{I("Y")})'
            f'&IF({I("Z")}="",""," · Área: "&{I("Z")})'
            f'&IF({I("AD")}="",""," · Req.: "&{I("AD")})'
            f'&IF(OR(TRIM({I("C")})="",TRIM({I("C")})="-",UPPER(TRIM({I("C")}))=UPPER(TRIM({I("G")}))),""," · Função solicitada: "&{I("C")})')
    cell(f'I{r}', f'=IF($L{r}="","",MID({tudo},4,2000))', size=8, wrap=True)
    cell(f'J{r}', f'=IF($L{r}="","",{I("D")})', h='center')
    ws[f'N{r}'] = f'=IF($L{r}="",1,"")'
    if k % 2 == 0:
        for c in 'BCDEFGHIJ': ws[f'{c}{r}'].fill = fill('FFF7F7F7')
    ws.row_dimensions[r].height = 24; r += 1
ws.conditional_formatting.add(f'B{e0}:J{r-1}', FormulaRule(formula=[f'ISNUMBER(SEARCH("⚠",$I{e0}))'], fill=fill(VERM)))
ws[f'B{r}'] = f'=IF(SUM(Lanc_HE!$AF$2:$AF$693)>{NHE},"ATENÇÃO: o mês tem mais de {NHE} HE; aumente o relatório.","")'
ws[f'N{r}'] = f'=IF(B{r}="",1,"")'; r += 1

# ---------- COMPENSAÇÕES ----------
NCP = 30
hidCP = f'=IF(SUM(Compensacoes!$N$4:$N$503)=0,1,"")'
r += 1; bar(r, 'COMPENSAÇÕES DO MÊS', hide=hidCP); r += 1
hdrrow(r, [('B', 'Data'), ('C', 'Quem trabalhou'), ('D', 'Compensa a falta de'), ('E', 'Horas'), ('F', 'Aprovado?'), ('G', 'Data da falta'),
           ('H', 'Horas que contam'), ('I', 'Situação / observação'), ('J', 'Equipe')], hide=hidCP); r += 1
c0 = r
for k in range(1, NCP + 1):
    ws[f'M{r}'] = k; ws[f'L{r}'] = f'=IFERROR(MATCH(M{r},Compensacoes!$O$4:$O$503,0)+3,"")'
    I = lambda c: f'INDEX(Compensacoes!${c}$1:${c}$503,$L{r})'
    cell(f'B{r}', f'=IF($L{r}="","",{I("A")})', h='center', fmt='dd/mm/yyyy'); cell(f'C{r}', f'=IF($L{r}="","",{I("B")})', wrap=True)
    cell(f'D{r}', f'=IF($L{r}="","",IF({I("D")}="","—",{I("D")}))', size=8, wrap=True)
    cell(f'E{r}', f'=IF($L{r}="","",N({I("C")}))', h='center', fmt='[h]:mm;;;')
    cell(f'F{r}', f'=IF($L{r}="","",{I("F")})', h='center')
    cell(f'G{r}', f'=IF($L{r}="","",IF({I("E")}="","",{I("E")}))', h='center', fmt='dd/mm/yyyy')
    cell(f'H{r}', f'=IF($L{r}="","",N({I("J")}))', h='center', fmt='[h]:mm;;;')
    cell(f'I{r}', f'=IF($L{r}="","",{I("L")}&IF({I("G")}="",""," · "&{I("G")}))', size=8, wrap=True)
    cell(f'J{r}', f'=IF($L{r}="","",{I("I")})', h='center')
    ws[f'N{r}'] = f'=IF($L{r}="",1,"")'; ws.row_dimensions[r].height = 20; r += 1
ws.conditional_formatting.add(f'F{c0}:F{r-1}', FormulaRule(formula=[f'AND($F{c0}<>"",$F{c0}<>"Sim")'], fill=fill(VERM)))
ws.conditional_formatting.add(f'F{c0}:F{r-1}', FormulaRule(formula=[f'$F{c0}="Sim"'], fill=fill(VERDE)))
r += 1
note(r, 'Fonte: planilha Lançamento de Horas EI32 (abas Lanc_HORAS, Lanc_HE, Compensações e Analises). Horas em h:mm; valores em R$. '
        'Linhas e grupos sem ocorrência no mês não aparecem no PDF.')
LAST = r
ws['L1'] = LAST
CU = wb['Como_Usar']
CU['C35'].value = ('Gera o PDF do mês escolhido na Gestão (B3): pendências para fechar o mês, ferista, todas as ocorrências por tipo '
                   '(descontam / justificadas / ajustes e trocas de turno / a conferir), o detalhe dia a dia com cobertura do ferista e motivo, '
                   'as horas extras por tratamento (efeito no contrato) e caso a caso, e as compensações.')
CU.row_dimensions[35].height = 40
ws.print_area = f'B1:J{LAST}'
wb.save(OUT)
print('LAST', LAST)
