import openpyxl,datetime
def num(v):
    if isinstance(v,datetime.timedelta): return round(v.total_seconds()/60,1)
    if isinstance(v,datetime.time): return v.hour*60+v.minute
    if isinstance(v,float): return round(v,4)
    return v
a=openpyxl.load_workbook('in/atual.xlsx',data_only=True,read_only=True); b=openpyxl.load_workbook('v53c.xlsx',data_only=True,read_only=True)
for s in ['Gestao_Contrato','Analises','Painel_De_Controle','SLA','Lanc_HE','Compensacoes']:
    A=[[num(c) for c in r] for r in a[s].iter_rows(values_only=True)]
    B=[[num(c) for c in r] for r in b[s].iter_rows(values_only=True)]
    d=[(i+1,j+1,A[i][j],B[i][j]) for i in range(min(len(A),len(B))) for j in range(min(len(A[i]),len(B[i]))) if A[i][j]!=B[i][j] and not (isinstance(A[i][j],(int,float)) and isinstance(B[i][j],(int,float)) and abs(A[i][j]-B[i][j])<0.01)]
    print(s,len(d),d[:6])
R1=list(a['Relatorio'].iter_rows(values_only=True)); R2=list(b['Relatorio'].iter_rows(values_only=True))
print('Rel B2',R1[1][1],'|',R2[1][1])
cnt=lambda R: sum(1 for r in R if r and isinstance(r[1],(datetime.datetime,)) )
print('linhas de detalhe com data: antes',cnt(R1),'depois',cnt(R2))
print('ocorr resumo',R1[6][4],R1[7][4],'|',R2[6][4],R2[7][4])
