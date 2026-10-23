"""v55 — Abatimento de Valor/Função (ferista de cargo menor cobrindo técnico de cargo maior).

Regra (Lucas, 22/10/2026):
- o ferista abate as horas de absenteísmo dos técnicos (nunca do encarregado);
- se o coberto tem valor-hora maior que o do ferista, as horas cobertas NÃO geram absenteísmo,
  mas geram um débito = (valor HH coberto − valor HH ferista) × horas cobertas: "Abatimento de Valor/Função";
- ferista de cargo maior que o coberto: nada (nunca gera crédito);
- vale desde janeiro/2026 (parâmetro Tabelas!K18);
- Rodrigo Chanchencow Soares (544631) é TÉC. ELETRICISTA I (correção de cadastro/vigências).
"""
import openpyxl, copy, re, sys
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from datetime import datetime

SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)
lh, an, g, p, rel, tb = (wb[n] for n in ('Lanc_HORAS', 'Analises', 'Gestao_Contrato', 'Painel_De_Controle', 'Relatorio', 'Tabelas'))
NL = 3873
assert lh.max_row == NL, lh.max_row
TK, NAVY, TXT, CINZA = 'FF0095D8', 'FF1F3864', 'FF4D4D4D', 'FF7F7F7F'
st = lambda dst, src: setattr(dst, '_style', copy.copy(src._style))

# ------------------------------------------------------------ 1. Rodrigo = TÉC. ELETRICISTA I
RG_ROD = 544631
cad = wb['Cadastro']
for r in range(3, 151):
    if cad.cell(r, 1).value == RG_ROD:
        cad.cell(r, 5).value = 'TÉC. ELETRICISTA I'
        cad.cell(r, 11).value = 'Ferista contratado como Téc. Eletricista I (cadastro corrigido em 22/10/2026 — compliance).'
vig = wb['Vigencias']
# valores HH do Téc. Eletricista I em cada vigência (iguais aos dos outros Téc. I)
hh_tec1 = {}
for r in range(3, 150):
    if vig.cell(r, 5).value == 'TÉC. ELETRICISTA I' and isinstance(vig.cell(r, 6).value, (int, float)):
        hh_tec1[vig.cell(r, 4).value] = vig.cell(r, 6).value
ini_jan, ini_ago = datetime(2026, 1, 1), datetime(2026, 8, 1)
for r in range(3, 150):
    if vig.cell(r, 1).value == RG_ROD:
        d = vig.cell(r, 4).value
        vig.cell(r, 5).value = 'TÉC. ELETRICISTA I'
        vig.cell(r, 6).value = hh_tec1[ini_ago] if d >= ini_ago else hh_tec1[ini_jan]

# ------------------------------------------------------------ 2. Parâmetro de início
for c, v in ((10, 'Início do Abatimento de Valor/Função'), (11, datetime(2026, 1, 1)),
             (12, 'Ferista de cargo menor cobrindo técnico de cargo maior: desconta a diferença de valor-hora × horas cobertas a partir desta data.')):
    x = tb.cell(18, c); x.value = v; st(x, tb.cell(15, c))
tb['K18'].number_format = 'dd/mm/yyyy'

# ------------------------------------------------------------ 3. Lanc_HORAS
# AC: ferista nunca cobre encarregado (nem na cobertura manual)
for r in range(2, NL + 1):
    f = lh[f'AC{r}'].value
    old = f'=IF($A{r}="",0,'
    assert f.startswith(old), (r, f[:40])
    lh[f'AC{r}'] = f'=IF(OR($A{r}="",ISNUMBER(SEARCH("ENCARREGADO",$G{r}))),0,' + f[len(old):]

VIG = lambda col, rg, d: (f'LOOKUP(2,1/((Vigencias!$A$3:$A$149={rg})*(Vigencias!$D$3:$D$149<={d})),Vigencias!${col}$3:${col}$149)')
heads = {'AY': 'RG do ferista que cobriu (aux)', 'AZ': 'Valor HH do ferista (aux)',
         'BA': 'Abatimento\nValor/Função (R$)', 'BB': 'Abatimento Valor/Função —\ncomo foi calculado'}
for c, h in heads.items():
    lh[f'{c}1'] = h; st(lh[f'{c}1'], lh['AX1'])
rng = lambda c: f'${c}$2:${c}${NL}'
for r in range(2, NL + 1):
    lh[f'AY{r}'] = (f'=IF(N($AC{r})<=0,"",IF(COUNTIFS({rng("A")},$A{r},{rng("D")},$D{r},{rng("C")},"Ferista")=1,'
                    f'SUMIFS({rng("F")},{rng("A")},$A{r},{rng("D")},$D{r},{rng("C")},"Ferista"),'
                    f'IFERROR(LOOKUP(2,1/((Vigencias!$C$3:$C$149=$D{r})*(Vigencias!$H$3:$H$149="Ferista")*(Vigencias!$D$3:$D$149<=$A{r})),Vigencias!$A$3:$A$149),"")))')
    lh[f'AZ{r}'] = (f'=IF($AY{r}="","",IFERROR({VIG("F", f"$AY{r}", f"$A{r}")},'
                    f'IFERROR(INDEX(Cadastro!$G$3:$G$150,MATCH($AY{r},Cadastro!$A$3:$A$150,0)),"")))')
    lh[f'BA{r}'] = (f'=IF(OR($A{r}="",$AZ{r}="",N($AC{r})<=0,NOT(ISNUMBER($I{r})),$C{r}="Ferista",'
                    f'ISNUMBER(SEARCH("ENCARREGADO",$G{r}))),0,IF($A{r}<N(Tabelas!$K$18),0,'
                    f'ROUND(MAX(ROUND($I{r},2)-ROUND($AZ{r},2),0)*N($AC{r})*24,2)))')
    lh[f'BB{r}'] = (f'=IF(N($BA{r})<=0,"","Acobertado pelo ferista / Abatimento de Valor/Função: ferista "'
                    f'&IFERROR({VIG("E", f"$AY{r}", f"$A{r}")},"")&" (R$ "&FIXED($AZ{r},2)&"/h) cobrindo "&$G{r}'
                    f'&" (R$ "&FIXED($I{r},2)&"/h). Diferença R$ "&FIXED(ROUND($I{r},2)-ROUND($AZ{r},2),2)&"/h × "'
                    f'&TEXT($AC{r},"[h]:mm")&"h cobertas = R$ "&FIXED($BA{r},2)&". Não gera absenteísmo.")')
    for c, src in (('AY', 'AU'), ('AZ', 'AU'), ('BA', 'W'), ('BB', 'AA')):
        st(lh[f'{c}{r}'], lh[f'{src}{r}'])
    lh[f'BA{r}'].number_format = '#,##0.00;;;'
for c in ('AY', 'AZ'):
    lh.column_dimensions[c].hidden = True
lh.column_dimensions['BA'].width = 14
lh.column_dimensions['BB'].width = 70

# ------------------------------------------------------------ 4. Analises
for b0, team in ((3, 'Elétrica'), (19, 'Climatização')):
    hr = b0 - 1
    an[f'BH{hr}'] = 'Abatimento\nValor/Função (R$)'; an[f'BI{hr}'] = 'Horas cobertas com\ndiferença de função (h)'
    st(an[f'BH{hr}'], an[f'BG{hr}']); st(an[f'BI{hr}'], an[f'BG{hr}'])
    for r in range(b0, b0 + 12):
        an[f'BH{r}'] = f'=SUMIFS(Lanc_HORAS!$BA$2:$BA${NL},Lanc_HORAS!$B$2:$B${NL},$A{r},Lanc_HORAS!$D$2:$D${NL},"{team}")'
        an[f'BI{r}'] = (f'=SUMIFS(Lanc_HORAS!$AC$2:$AC${NL},Lanc_HORAS!$B$2:$B${NL},$A{r},Lanc_HORAS!$D$2:$D${NL},"{team}",'
                        f'Lanc_HORAS!$BA$2:$BA${NL},">0")')
        st(an[f'BH{r}'], an[f'I{r}']); st(an[f'BI{r}'], an[f'BG{r}'])
        assert an[f'I{r}'].value == f'=MAX(N($AV{r})-N($AT{r}),0)', an[f'I{r}'].value
        an[f'I{r}'] = f'=MAX(N($AV{r})-N($AT{r}),0)+N($BH{r})'
    tr = b0 + 12
    an[f'BH{tr}'] = f'=SUM(BH{b0}:BH{tr-1})'; an[f'BI{tr}'] = f'=SUM(BI{b0}:BI{tr-1})'
    st(an[f'BH{tr}'], an[f'I{tr}']); st(an[f'BI{tr}'], an[f'BG{tr}'])
an.column_dimensions['BH'].width = 16; an.column_dimensions['BI'].width = 18

# ------------------------------------------------------------ 5. Gestão: nova linha 21
REF = re.compile(r"((?:'[^']+'|[A-Za-z_][A-Za-z0-9_\.]*)!)?(\$?)([A-Z]{1,3})(\$?)(\d+)(?::(\$?)([A-Z]{1,3})(\$?)(\d+))?")
def shift_local(f, ini, n=1):
    def rep(m):
        if m.group(1):
            return m.group(0)
        a = int(m.group(5)); s = m.group(2) + m.group(3) + m.group(4) + str(a + n if a >= ini else a)
        if m.group(7):
            b = int(m.group(9)); s += ':' + m.group(6) + m.group(7) + m.group(8) + str(b + n if b >= ini else b)
        return s
    out, i = [], 0
    for q in re.finditer(r'"[^"]*"', f):   # não mexe dentro de textos
        out.append(REF.sub(rep, f[i:q.start()])); out.append(q.group(0)); i = q.end()
    out.append(REF.sub(rep, f[i:]))
    return ''.join(out)

INS = 21
last = g.max_row
merges = [m for m in g.merged_cells.ranges if m.min_row >= INS]
for m in merges:
    g.unmerge_cells(str(m))
for r in range(last, INS - 1, -1):
    for c in range(1, g.max_column + 1):
        s, d = g.cell(r, c), g.cell(r + 1, c)
        v = s.value
        d.value = shift_local(v, INS) if isinstance(v, str) and v.startswith('=') else v
        st(d, s)
    g.row_dimensions[r + 1].height = g.row_dimensions[r].height
for m in merges:
    g.merge_cells(start_row=m.min_row + 1, start_column=m.min_col, end_row=m.max_row + 1, end_column=m.max_col)
IXG = lambda col, b0: f'IFERROR(INDEX(Analises!${col}${b0}:${col}${b0+11},MATCH($B$3,Analises!$A${b0}:$A${b0+11},0)),0)'
for (lb, hc, vc, b0) in (('B', 'C', 'D', 3), ('F', 'G', 'H', 19)):
    for c in (lb, hc, vc):
        st(g[f'{c}{INS}'], g[f'{c}20'])
    g[f'{lb}{INS}'] = '( + ) Abatimento de Valor/Função (ferista de cargo menor que o coberto)'
    g[f'{hc}{INS}'] = '=' + IXG('BI', b0)
    g[f'{vc}{INS}'] = '=' + IXG('BH', b0)
    g[f'{lb}{INS}'].font = Font(name='Arial', size=9, bold=True, color='FFC00000')
    g[f'{vc}{INS}'].font = Font(name='Arial', size=9, bold=True, color='FFC00000')
    # compensação passa a ser o valor compensado (antes era AV − I, que agora incluiria o valor/função)
    g[f'{vc}20'] = '=' + IXG('AT', b0)
    tot = INS + 1
    assert g[f'{lb}{tot}'].value == 'TOTAL ABATIDO'
    g[f'{vc}{tot}'] = f'=MAX(SUM({vc}11:{vc}19)-{vc}20,0)+{vc}{INS}'
g.row_dimensions[INS].height = g.row_dimensions[20].height

# ------------------------------------------------------------ 6. Painel de Controle
p['E40'] = 'HORAS PENDENTES'
for lb, vc, b0 in (('B', 'C', 3), ('E', 'F', 19)):
    st(p[f'{lb}46'], p[f'{lb}37']); st(p[f'{vc}46'], p[f'{vc}37'])
    p[f'{lb}46'] = 'ABATIMENTO DE VALOR/FUNÇÃO'
    p.merge_cells(f'{lb}46:{vc}46')
    for r, txt, col, fmt, srow in ((47, 'Horas cobertas por ferista de cargo menor', 'BI', '[h]:mm', 38),
                                   (48, 'Valor abatido — diferença de função (R$)', 'BH', '#,##0.00', 39)):
        st(p[f'{lb}{r}'], p[f'{lb}{srow}']); st(p[f'{vc}{r}'], p[f'{vc}{srow}'])
        p[f'{lb}{r}'] = txt
        p[f'{vc}{r}'] = f'=IFERROR(INDEX(Analises!${col}${b0}:${col}${b0+11},MATCH($C$2,Analises!$A${b0}:$A${b0+11},0)),0)'
        p[f'{vc}{r}'].number_format = fmt
    p[f'{vc}48'].font = Font(name='Arial', size=9, bold=True, color='FFC00000')

# ------------------------------------------------------------ 7. Relatório
def acha(texto):
    return [c.row for c in rel['C'] if c.value == texto]
for abat_row, b0 in zip(acha('ABATIMENTO DO MÊS (resultado)'), (3, 19)):
    comp = abat_row - 1; new = abat_row; dst = abat_row + 1
    assert rel[f'C{comp}'].value.startswith('( − ) Compensações') and all(rel.cell(dst, c).value is None for c in range(2, 15))
    for m in [m for m in rel.merged_cells.ranges if m.min_row == new]:
        rel.unmerge_cells(str(m))
    for c in range(2, 15):
        s, d = rel.cell(new, c), rel.cell(dst, c)
        d.value = s.value; st(d, s)
        s.value = None; st(s, rel.cell(comp, c))
    rel.merge_cells(f'C{dst}:D{dst}'); rel.merge_cells(f'I{dst}:J{dst}')
    rel.row_dimensions[dst].height = rel.row_dimensions[new].height
    rel[f'C{new}'] = '( + ) Abatimento de Valor/Função (ferista de cargo menor que o coberto)'
    rel[f'G{new}'] = f'=IFERROR(INDEX(Analises!$BI${b0}:$BI${b0+11},MATCH(Gestao_Contrato!$B$3,Analises!$A${b0}:$A${b0+11},0)),0)'
    rel[f'H{new}'] = f'=IFERROR(INDEX(Analises!$BH${b0}:$BH${b0+11},MATCH(Gestao_Contrato!$B$3,Analises!$A${b0}:$A${b0+11},0)),0)'
    rel[f'I{new}'] = ('Ferista cobriu técnico de função mais cara: as horas não geram absenteísmo, '
                      'mas desconta a diferença de valor-hora × horas cobertas.')
    rel[f'N{new}'] = f'=IF(N(H{new})=0,1,"")'
    rel.merge_cells(f'C{new}:D{new}'); rel.merge_cells(f'I{new}:J{new}')
    rel[f'C{new}'].font = Font(name='Arial', size=9, bold=True, color='FFC00000')
    rel[f'H{new}'].font = Font(name='Arial', size=9, bold=True, color='FFC00000')
    rel.row_dimensions[new].height = 24

# FERISTA DO MÊS: leitura mostra o valor/função quando houver
for r, b0 in ((20, 3), (21, 19)):
    base = rel[f'I{r}'].value
    assert isinstance(base, str) and base.startswith('Cobriu =')
    v = f'IFERROR(INDEX(Analises!$BH${b0}:$BH${b0+11},MATCH(Gestao_Contrato!$B$3,Analises!$A${b0}:$A${b0+11},0)),0)'
    rel[f'I{r}'] = (f'=IF({v}>0,"Abatimento de Valor/Função: R$ "&FIXED({v},2)&" (ferista de cargo menor cobrindo técnico de cargo maior). ","")'
                    f'&"{base}"')
    rel.row_dimensions[r].height = 30

# Detalhe dia a dia
IX = lambda c, r: f'INDEX(Lanc_HORAS!${c}$1:${c}${NL},$L{r})'
n_det = 0
for r in range(1, rel.max_row + 1):
    lv = rel[f'L{r}'].value
    if not (isinstance(lv, str) and lv.startswith('=IFERROR(MATCH(M') and 'Lanc_HORAS!' in lv):
        continue
    n_det += 1
    ba = f'N({IX("BA", r)})'
    d = rel[f'D{r}'].value; assert d.startswith(f'=IF($L{r}="","",') and d.endswith(')')
    rel[f'D{r}'] = d[:-1] + f'&IF({ba}>0," · Acobertado pelo ferista / Abatimento de Valor/Função",""))'
    f = rel[f'F{r}'].value; assert f.startswith(f'=IF($L{r}="","",') and f.endswith(')')
    rel[f'F{r}'] = f[:-1] + f'&IF({ba}>0," · Abat. Valor/Função R$ "&FIXED({ba},2),""))'
    h = rel[f'H{r}'].value; assert h == f'=IF($L{r}="","",N({IX("W", r)}))', h
    rel[f'H{r}'] = f'=IF($L{r}="","",N({IX("W", r)})+{ba})'
    i = rel[f'I{r}'].value; assert i == f'=IF($L{r}="","",{IX("AA", r)})', i
    rel[f'I{r}'] = f'=IF($L{r}="","",TRIM({IX("AA", r)}&" "&{IX("BB", r)}))'
assert n_det == 300, n_det

# ------------------------------------------------------------ 8. Legenda
lg = wb['Legenda']
r0 = lg.max_row + 2
st(lg[f'B{r0}'], lg['B51']); st(lg[f'C{r0}'], lg['C51'])
lg[f'B{r0}'] = 'ABATIMENTO DE VALOR/FUNÇÃO (ferista)'; lg.merge_cells(f'B{r0}:C{r0}')
itens = [('Regra', 'O ferista cobre faltas, atrasos, férias e remanejamentos dos TÉCNICOS da equipe (nunca do encarregado). '
                   'As horas cobertas não geram absenteísmo.'),
         ('Quando desconta', 'Se o técnico coberto tem valor-hora MAIOR que o do ferista (ex.: ferista Téc. I cobrindo Téc. II), '
                             'desconta a diferença de valor-hora × horas cobertas. Ferista de cargo igual ou maior: não desconta nada.'),
         ('Onde aparece', 'Lanc_HORAS colunas BA (valor) e BB (cálculo); Gestão linha "( + ) Abatimento de Valor/Função"; '
                          'Painel de Controle; Relatório (linha própria antes do ABATIMENTO DO MÊS e no detalhe dia a dia).'),
         ('Data de início', 'Tabelas › Parâmetros › "Início do Abatimento de Valor/Função" (hoje 01/01/2026).')]
for k, (a, b) in enumerate(itens, 1):
    st(lg[f'B{r0+k}'], lg['B52']); st(lg[f'C{r0+k}'], lg['C52'])
    lg[f'B{r0+k}'] = a; lg[f'C{r0+k}'] = b; lg.row_dimensions[r0 + k].height = 32

wb.save(OUT)
print('ok', n_det)
