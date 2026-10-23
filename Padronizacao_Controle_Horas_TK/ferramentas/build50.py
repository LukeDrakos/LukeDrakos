import openpyxl, re, copy, sys
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import Rule

SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)
NAVY, AZUL, BRANCO = 'FF1F3864', 'FF0095D8', 'FFFFFFFF'

# ---------- 1. renomear Resumo_Mensal -> Analises em todo o arquivo ----------
OLD, NEW = 'Resumo_Mensal', 'Analises'
def ren(s):
    return s.replace("'%s'!" % OLD, NEW + '!').replace(OLD + '!', NEW + '!')
R = wb[OLD]; R.title = NEW
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('=') and OLD in c.value:
                c.value = ren(c.value)
            if c.hyperlink is not None and c.hyperlink.location and OLD in c.hyperlink.location:
                c.hyperlink.location = c.hyperlink.location.replace("'%s'" % OLD, "'%s'" % NEW).replace(OLD, NEW)
    for cf in ws.conditional_formatting:
        for r in cf.rules:
            r.formula = [ren(f) for f in (r.formula or [])]
    for dv in ws.data_validations.dataValidation:
        if dv.formula1 and OLD in dv.formula1: dv.formula1 = ren(dv.formula1)
dn = wb.defined_names['lstMeses']; dn.attr_text = ren(dn.attr_text)
for n, d in wb.defined_names.items():
    assert OLD not in (d.attr_text or ''), n

# ---------- 2. colunas do Analises: expandir faixas e limpar agrupamento ----------
widths = {}
for k, d in list(R.column_dimensions.items()):
    lo, hi = (d.min or CI(k)), (d.max or CI(k))
    for i in range(lo, hi + 1):
        widths[i] = d.width
R.column_dimensions.clear()
def setw(i, w):
    R.column_dimensions[L(i)] = ColumnDimension(R, index=L(i), width=w)
for i, w in widths.items():
    if i <= 60 and w: setw(i, w)

# ---------- 3. mover blocos para o lado direito ----------
ref_re = re.compile(r'^(\$?)([A-Z]{1,3})(\$?)(\d+)$')
def shift_formula(f, dr, dc, minrow, src_title):
    tok = Tokenizer(f)
    out = []
    for t in tok.items:
        v = t.value
        if t.type == 'OPERAND' and t.subtype == 'RANGE' and '!' not in v:
            parts = v.split(':')
            if all(ref_re.match(p) for p in parts):
                newp = []
                for p in parts:
                    a, col, b, row = ref_re.match(p).groups()
                    row = int(row)
                    if row < minrow:
                        raise ValueError('ref fora do bloco %s em %s: %s' % (p, src_title, f))
                    newp.append('%s%s%s%d' % (a, L(CI(col) + dc), b, row + dr))
                v = ':'.join(newp)
        out.append(v)
    return '=' + ''.join(out)

def fix_filter(f):
    # $B$2 da aba de origem = Gestao_Contrato!$B$3
    return re.sub(r'(?<![A-Za-z_!$])\$B\$2(?!\d)', 'Gestao_Contrato!$B$3', f)

C0 = CI('L')
heights = {}
def bump(r, h):
    if h: heights[r] = max(heights.get(r, 0), h)

def copiar_bloco(nome, r_ini, r_fim, dest_row, dc, ncols):
    S = wb[nome]
    dr = dest_row - r_ini
    for r in range(r_ini, r_fim + 1):
        bump(r + dr, S.row_dimensions[r].height or 15)
        for c in range(1, ncols + 1):
            sc = S.cell(r, c); dcell = R.cell(r + dr, c + dc)
            v = sc.value
            if isinstance(v, str) and v.startswith('='):
                v = shift_formula(fix_filter(v), dr, dc, r_ini, nome)
            dcell.value = v
            if sc.has_style: dcell._style = copy.copy(sc._style)
    for m in S.merged_cells.ranges:
        if m.min_row >= r_ini and m.max_row <= r_fim:
            R.merge_cells(start_row=m.min_row + dr, end_row=m.max_row + dr,
                          start_column=m.min_col + dc, end_column=m.max_col + dc)
    for cf in S.conditional_formatting:
        rngs = []
        for cr in cf.sqref.ranges:
            rngs.append('%s%d:%s%d' % (L(cr.min_col + dc), cr.min_row + dr, L(cr.max_col + dc), cr.max_row + dr))
        for rule in cf.rules:
            nr = copy.copy(rule)
            nr.formula = [shift_formula('=' + f, dr, dc, r_ini, nome)[1:] for f in rule.formula]
            R.conditional_formatting.add(' '.join(rngs), nr)
    for c in range(1, ncols + 1):
        w = S.column_dimensions[L(c)].width or 9
        i = c + dc
        widths_dir[i] = max(widths_dir.get(i, 0), w)

widths_dir = {}
def titulo_bloco(row, texto, anchor_name):
    R.merge_cells(start_row=row, end_row=row, start_column=C0, end_column=C0 + 29)
    c = R.cell(row, C0, texto)
    c.font = Font(name='Arial', size=12, bold=True, color=BRANCO)
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    bump(row, 22)
    for k in range(C0 + 1, C0 + 30):
        R.cell(row, k).fill = PatternFill('solid', fgColor=NAVY)

anchors = {}
row = 37
anchors['col'] = row; titulo_bloco(row, 'HORAS POR COLABORADOR  (mês do filtro)', 'col'); row += 1
copiar_bloco('Horas_por_Colaborador', 4, 35, row, C0 - 1, 30); row += 32 + 2
anchors['fun'] = row; titulo_bloco(row, 'HORAS POR FUNÇÃO  (mês do filtro)', 'fun'); row += 1
copiar_bloco('Horas_por_Funcao', 4, 20, row, C0, 29); row += 17 + 2
anchors['fer'] = row; titulo_bloco(row, 'COBERTURA DO FERISTA  (mês do filtro)', 'fer'); row += 1
copiar_bloco('Ferista', 4, 7, row, C0, 7); row += 4
FIM = row - 1


# preencher coluna L nos blocos Função/Ferista (alinhados a partir de M) e alinhar títulos de seção
for r in range(anchors['fun'] + 1, FIM + 1):
    if r in (anchors['fer'] - 1, anchors['fer'] - 2, anchors['fer']):
        continue
    m = R.cell(r, C0 + 1)
    R.cell(r, C0)._style = copy.copy(m._style)
for mr in list(R.merged_cells.ranges):
    if mr.min_row > anchors['fun'] and mr.min_col == C0 + 1:
        a = (mr.min_row, mr.max_row, mr.max_col)
        R.unmerge_cells(str(mr))
        v = R.cell(a[0], C0 + 1).value
        R.cell(a[0], C0 + 1).value = None
        R.cell(a[0], C0).value = v
        R.merge_cells(start_row=a[0], end_row=a[1], start_column=C0, end_column=a[2])
for r in range(37, FIM + 1):
    for col in (C0, C0 + 1):
        c = R.cell(r, col)
        if c.value in ('Elétrica', 'Climatização'):
            c.alignment = Alignment(horizontal='left', vertical='center', indent=1)

# larguras: coluna K = respiro; L.. = max dos blocos
setw(CI('K'), 3)
for i, w in widths_dir.items():
    setw(i, min(w, 62))
for r, h in heights.items():
    R.row_dimensions[r].height = max(R.row_dimensions[r].height or 0, h)

# ---------- 4. cabeçalho e índice de navegação da aba ----------
R['A33'].value = 'EI32 — ANÁLISES DO MÊS   ◄ Gestão'
R['A33'].hyperlink = None
R['A33'].value = '=HYPERLINK("#Gestao_Contrato!B3","EI32 — ANÁLISES DO MÊS   ◄ Gestão")'
R.row_dimensions[34].height = 20
nav = [('A34', 'Ir para:', None),
       ('C34', '► Resumo do mês (cartões)', 'B37'),
       ('H34', '► Colaboradores', 'L%d' % anchors['col']),
       ('I34', '► Funções', 'L%d' % anchors['fun']),
       ('J34', '► Ferista', 'L%d' % anchors['fer'])]
for addr, txt, tgt in nav:
    c = R[addr]
    c.value = txt if tgt is None else '=HYPERLINK("#Analises!%s","%s")' % (tgt, txt)
    c.font = Font(name='Arial', size=10, bold=True, color=NAVY if tgt else 'FF4D4D4D', underline='single' if tgt else None)
    c.alignment = Alignment(horizontal='left', vertical='center')
R.merge_cells('A34:B34')
R.freeze_panes = 'A36'
R.sheet_view.topLeftCell = 'A32'
R.sheet_properties.tabColor = wb['Lanc_HORAS'].sheet_properties.tabColor or 'FF0095D8'

# ---------- 5. remover as abas antigas e reordenar ----------
for s in ['Horas_por_Colaborador', 'Horas_por_Funcao', 'Ferista']:
    wb.remove(wb[s])
order = ['Como_Usar', 'IT_Instrucao_Trabalho', 'Legenda', 'Gestao_Contrato', 'Painel_De_Controle', 'Relatorio',
         'Analises', 'Lanc_HORAS', 'Lanc_HE', 'Compensacoes', 'Conferencia_Ponto', 'Cadastro', 'Vigencias', 'Tabelas', 'Painel_Aux']
assert sorted(order) == sorted(wb.sheetnames), wb.sheetnames
wb._sheets = [wb[n] for n in order]
wb.active = order.index('Gestao_Contrato')
for ws in wb.worksheets:
    ws.sheet_view.tabSelected = (ws.title == 'Gestao_Contrato')

# ---------- 6. menu da Gestão ----------
G = wb['Gestao_Contrato']
menu = {'C2': ('Análises (resumo do mês)', "'Analises'!B37"),
        'D2': ('Colaboradores', "'Analises'!L%d" % anchors['col']),
        'E2': ('Funções', "'Analises'!L%d" % anchors['fun']),
        'G2': ('Ferista', "'Analises'!L%d" % anchors['fer'])}
for a, (t, loc) in menu.items():
    G[a].value = t
    G[a].hyperlink = None
    G[a].hyperlink = openpyxl.worksheet.hyperlink.Hyperlink(ref=a, location=loc, display=t)
G['C3'].alignment = Alignment(horizontal=G['C3'].alignment.horizontal, vertical='center', shrink_to_fit=True)

# ---------- 7. liberar células para os botões ----------
REL = wb['Relatorio']
REL.unmerge_cells('B1:J1'); REL.merge_cells('B1:H1')
CON = wb['Conferencia_Ponto']
CON.unmerge_cells('B1:J1'); CON.merge_cells('B1:F1')
for ws, rng in [(REL, 'I1:J1'), (CON, 'G1:J1')]:
    for row_ in ws[rng]:
        for c in row_:
            c.value = None
            c._style = copy.copy(ws['B1']._style)

# ---------- 8. textos ----------
PA = wb['Painel_Aux']
for a in ['I1', 'J1', 'M1', 'N1']:
    PA[a].value = PA[a].value.replace('Horas_por_Colaborador', 'Analises/Colaboradores').replace('Horas_por_Funcao', 'Analises/Funções')
CU = wb['Como_Usar']
CU['C24'].value = 'Gestao_Contrato no mês certo (B3) e conferida com a aba Analises.'
for col in ('B', 'C'):
    CU[col + '12']._style = copy.copy(CU[col + '11']._style)
CU.row_dimensions[12].height = 27.75
CU['B12'].value = 'Analises'
CU['C12'].value = ('Resumo do mês em cartões (à esquerda) e, à direita, horas por colaborador, por função e do ferista. '
                   'Atalhos na linha 34 e no menu da Gestão.')
CU['B12'].hyperlink = openpyxl.worksheet.hyperlink.Hyperlink(ref='B12', location="'Analises'!A33", display='Analises')
LG = wb['Legenda']
if LG['B46'].value: LG['B46'].value = LG['B46'].value.replace('Resumo Mensal', 'Análises — Resumo do mês')

CU['B27'].hyperlink = openpyxl.worksheet.hyperlink.Hyperlink(ref='B27', location="'Analises'!L%d" % anchors['fer'], display='Ferista')
wb.save(OUT)
print('ok', anchors, 'fim', FIM)
