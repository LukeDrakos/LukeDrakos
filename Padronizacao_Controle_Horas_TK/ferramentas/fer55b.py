"""v55b — Gestão: tira as horas da linha do Abatimento de Valor/Função e cria o quadro
"FERISTA × CARGO COBERTO — RELAÇÃO DO VALOR-HORA" (só valores por hora, sem relação de horas)."""
import openpyxl, copy, sys
from openpyxl.styles import Font, Alignment, PatternFill

SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)
g, p, rel, cad = wb['Gestao_Contrato'], wb['Painel_De_Controle'], wb['Relatorio'], wb['Cadastro']
st = lambda dst, src: setattr(dst, '_style', copy.copy(src._style))
VERM = 'FFC00000'

# 1. linha do abatimento sem horas (Gestão e Relatório)
for c in ('C21', 'G21'):
    assert 'BI' in g[c].value
    g[c] = None
for r in [c.row for c in rel['C'] if c.value and str(c.value).startswith('( + ) Abatimento de Valor')]:
    rel[f'G{r}'] = None

# 2. quadro novo na Gestão
MES = '$B$3'
FIM = 'EOMONTH($B$3,0)'
def funcoes(team):
    fs = []
    for r in range(3, 151):
        if cad.cell(r, 4).value == team and cad.cell(r, 3).value == 'Normal / Fixo':
            f = cad.cell(r, 5).value
            if f and 'ENCARREGADO' not in f.upper() and f not in fs:
                fs.append(f)
    return sorted(fs)

R0 = g.max_row + 2          # 40
blocos = (('B', 'C', 'D', 'Elétrica', 'ELÉTRICA', 95, 3), ('F', 'G', 'H', 'Climatização', 'CLIMATIZAÇÃO', 96, 19))
nmax = max(len(funcoes(t)) for _, _, _, t, *_ in blocos)
for lb, c1, c2, team, TEAM, arow, b0 in blocos:
    r = R0
    # cabeçalho
    for c in (lb, c1, c2): st(g[f'{c}{r}'], g[f'{lb}23'])
    g[f'{lb}{r}'] = f'FERISTA × CARGO COBERTO — RELAÇÃO DO VALOR-HORA ({TEAM})'
    g.merge_cells(f'{lb}{r}:{c2}{r}'); r += 1
    # ferista do mês
    nome = f'Analises!$N${arow}'
    rg = f'IFERROR(INDEX(Cadastro!$A$3:$A$150,MATCH({nome},Cadastro!$B$3:$B$150,0)),"")'
    fun = f'IFERROR(LOOKUP(2,1/((Vigencias!$A$3:$A$149={rg})*(Vigencias!$D$3:$D$149<={FIM})),Vigencias!$E$3:$E$149),"")'
    hh = f'IFERROR(LOOKUP(2,1/((Vigencias!$A$3:$A$149={rg})*(Vigencias!$D$3:$D$149<={FIM})),Vigencias!$F$3:$F$149),"")'
    for c in (lb, c1, c2): st(g[f'{c}{r}'], g[f'{c}25'] if c != c2 else g[f'{c1}25'])
    g[f'{lb}{r}'] = 'Ferista do mês'
    g[f'{c1}{r}'] = f'=IF(OR({nome}="",{nome}="—"),"—",{nome})'
    g.merge_cells(f'{c1}{r}:{c2}{r}'); r += 1
    for c in (lb, c1, c2): st(g[f'{c}{r}'], g[f'{c}26'] if c != c2 else g[f'{c1}26'])
    g[f'{lb}{r}'] = 'Função e valor-hora do ferista'
    g[f'{c1}{r}'] = f'={fun}'
    g[f'{c2}{r}'] = f'={hh}'
    g[f'{c2}{r}'].number_format = '"R$ "#,##0.00"/h"'
    fer_hh = f'${c2}${r}'
    r += 1
    # subcabeçalho
    for c, t in ((lb, 'Cargo coberto (técnico)'), (c1, 'Valor-hora do cargo'), (c2, 'Diferença para o ferista')):
        st(g[f'{c}{r}'], g[f'{lb}24' if c == lb else f'{c1}24']); g[f'{c}{r}'] = t
    r += 1
    fs = funcoes(team)
    for k in range(nmax):
        zeb = g[f'{lb}26'] if k % 2 else g[f'{lb}25']
        zv = g[f'{c1}26'] if k % 2 else g[f'{c1}25']
        st(g[f'{lb}{r}'], zeb); st(g[f'{c1}{r}'], zv); st(g[f'{c2}{r}'], zv)
        if k < len(fs):
            f = fs[k]
            g[f'{lb}{r}'] = f
            g[f'{c1}{r}'] = (f'=IFERROR(LOOKUP(2,1/((Vigencias!$E$3:$E$149={lb}{r})*(Vigencias!$D$3:$D$149<={FIM})),Vigencias!$F$3:$F$149),'
                             f'IFERROR(INDEX(Tabelas!$F$3:$F$150,MATCH({lb}{r},Tabelas!$E$3:$E$150,0)),""))')
            g[f'{c2}{r}'] = f'=IF(OR({fer_hh}="",{c1}{r}=""),"",MAX(ROUND({c1}{r},2)-ROUND({fer_hh},2),0))'
            g[f'{c1}{r}'].number_format = '"R$ "#,##0.00"/h"'
            g[f'{c2}{r}'].number_format = '"R$ "#,##0.00"/h";;"sem desconto"'
            g[f'{c2}{r}'].font = Font(name='Arial', size=9, bold=True, color=VERM)
        r += 1
    # total
    st(g[f'{lb}{r}'], g['B22']); st(g[f'{c1}{r}'], g['C22']); st(g[f'{c2}{r}'], g['D22'])
    g[f'{lb}{r}'] = 'ABATIMENTO DE VALOR/FUNÇÃO DO MÊS'
    g[f'{c1}{r}'] = None
    g[f'{c2}{r}'] = f'=IFERROR(INDEX(Analises!$BH${b0}:$BH${b0+11},MATCH($B$3,Analises!$A${b0}:$A${b0+11},0)),0)'
    g[f'{c2}{r}'].font = Font(name='Arial', size=g['D22'].font.sz or 11, bold=True, color=VERM)
    r += 1
    g[f'{lb}{r}'] = ('Só desconta quando o cargo coberto vale mais que o ferista; ferista igual ou maior = sem desconto. '
                     'Encarregado nunca é coberto. O cálculo de cada dia está na Lanc_HORAS (col. BB) e no Relatório.')
    g[f'{lb}{r}'].font = Font(name='Arial', size=8, italic=True, color='FF7F7F7F')
    g[f'{lb}{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    g.merge_cells(f'{lb}{r}:{c2}{r}')
    g.row_dimensions[r].height = 26
for rr in range(R0, R0 + nmax + 6):
    g.row_dimensions[rr].height = g.row_dimensions[rr].height or 15

# a linha 21 aponta para o quadro
for lb in ('B', 'F'):
    g[f'{lb}21'] = '( + ) Abatimento de Valor/Função (ver quadro abaixo)'

# 3. Painel: troca "horas cobertas" pelo valor-hora do ferista
for lb, vc, arow in (('B', 'C', 95), ('E', 'F', 96)):
    nome = f'Analises!$N${arow}'
    rg = f'IFERROR(INDEX(Cadastro!$A$3:$A$150,MATCH({nome},Cadastro!$B$3:$B$150,0)),"")'
    p[f'{lb}47'] = 'Valor-hora do ferista no mês (R$/h)'
    p[f'{vc}47'] = f'=IFERROR(LOOKUP(2,1/((Vigencias!$A$3:$A$149={rg})*(Vigencias!$D$3:$D$149<=EOMONTH($C$2,0))),Vigencias!$F$3:$F$149),"—")'
    p[f'{vc}47'].number_format = '#,##0.00'

wb.save(OUT)
print('ok', R0, nmax)
